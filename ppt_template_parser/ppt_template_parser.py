from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
import json
import os
from style_parser import ShapeStyleParser

class PPTTemplateParser:

    """
    PPT模板解析器

    输入:
        一个pptx文件

    输出:
        结构化JSON模板
    """


    def __init__(self, ppt_path):

        if not os.path.exists(ppt_path):
            raise FileNotFoundError(
                ppt_path
            )

        self.ppt_path = ppt_path

        self.prs = Presentation(
            ppt_path
        )

        # 新增：初始化样式解析器
        self.style_parser = ShapeStyleParser()


    # =========================
    # PPT尺寸
    # =========================

    def parse_slide_size(self):

        return {

            "width":
                self.prs.slide_width,

            "height":
                self.prs.slide_height

        }



    # =========================
    # 母版解析
    # =========================

    def parse_masters(self):

        masters=[]


        for master in self.prs.slide_masters:


            master_info={

                "name":
                    master.name,

                "layouts":[]

            }


            for layout in master.slide_layouts:

                master_info["layouts"].append(

                    self.parse_layout(layout)

                )


            masters.append(
                master_info
            )


        return masters



    # =========================
    # Layout解析
    # =========================

    def parse_layout(self,layout):


        layout_info={


            "name":
                layout.name,


            "placeholders":[]

        }


        for ph in layout.placeholders:


            item={


                "name":
                    ph.name,


                "type":
                    str(
                        ph.placeholder_format.type
                    ),


                "position":
                    self.get_position(ph)


            }


            layout_info[
                "placeholders"
            ].append(item)


        return layout_info



    # =========================
    # 所有Slide解析
    # =========================

    def parse_slides(self):


        slides=[]


        for index,slide in enumerate(
            self.prs.slides
        ):


            slide_info={


                "slide_id":
                    index+1,


                "layout":
                    slide.slide_layout.name,


                "elements":[]

            }


            for shape in slide.shapes:


                slide_info["elements"].append(

                    self.parse_shape(shape)

                )


            slides.append(
                slide_info
            )


        return slides



    # =========================
    # Shape解析
    # =========================

    def parse_shape(self, shape):


        info = {


            "name":
                shape.name,


            "type":
                str(
                    shape.shape_type
                ),


            "position":
                self.get_position(shape)

        }



        # =====================
        # 文本
        # =====================

        if shape.has_text_frame:


            info["text"] = shape.text


            info["text_style"] = (
                self.parse_text(shape)
            )



        # =====================
        # 图片
        # =====================

        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:


            info["image"] = {


                "width":
                    shape.image.size[0],


                "height":
                    shape.image.size[1]


            }




    # =========================
    # 坐标
    # =========================

    def get_position(self,shape):


        return {


            "left":
                shape.left,


            "top":
                shape.top,


            "width":
                shape.width,


            "height":
                shape.height


        }



    # =========================
    # 文本样式
    # =========================

    def parse_text(self,shape):


        styles=[]


        for p in shape.text_frame.paragraphs:


            for r in p.runs:


                styles.append({

                    "font":
                        r.font.name,


                    "size":
                        str(
                            r.font.size
                        ),


                    "bold":
                        r.font.bold,


                    "italic":
                        r.font.italic


                })


        return styles



    # =========================
    # 总入口
    # =========================

    def parse(self):


        result={


            "template_name":
                os.path.basename(
                    self.ppt_path
                ),


            "slide_size":
                self.parse_slide_size(),


            "masters":
                self.parse_masters(),


            "slides":
                self.parse_slides()


        }


        return result



    # 保存JSON

    def save_json(
        self,
        output_path
    ):


        data=self.parse()


        with open(
            output_path,
            "w",
            encoding="utf-8"
        ) as f:


            json.dump(

                data,

                f,

                indent=4,

                ensure_ascii=False

            )


        print(
            "保存完成:",
            output_path
        )
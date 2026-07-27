"""Small python-pptx helpers for chart, table, and image Visualization JSON."""

from __future__ import annotations

import math
import re
import unicodedata
from copy import deepcopy
from pathlib import Path
from typing import Any, Mapping, Sequence

from PIL import Image as PILImage
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import (
    XL_CHART_TYPE,
    XL_LABEL_POSITION,
    XL_LEGEND_POSITION,
    XL_MARKER_STYLE,
)
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Pt


class VisualizationRenderError(ValueError):
    """Raised for unsupported or inconsistent visualization data."""


DEFAULT_CHART_COLORS = (
    "2F75B5",
    "ED7D31",
    "70AD47",
    "A5A5A5",
    "FFC000",
    "5B9BD5",
)
DEFAULT_TEXT_COLOR = "334155"
DEFAULT_AXIS_COLOR = "94A3B8"
DEFAULT_GRIDLINE_COLOR = "E2E8F0"
DEFAULT_TABLE_HEADER_FILL = "102A43"
DEFAULT_TABLE_HEADER_TEXT = "FFFFFF"
DEFAULT_TABLE_BODY_FILL = "FFFFFF"
DEFAULT_TABLE_STRIPE_FILL = "F5F7FA"
DEFAULT_TABLE_BORDER = "D9E2EC"
DEFAULT_TABLE_TEXT = "334155"
DEFAULT_TABLE_MAX_ROWS = 18
DEFAULT_TABLE_MAX_COLUMNS = 8


def _chart_type(value: str):
    values = {
        "line": XL_CHART_TYPE.LINE_MARKERS,
        "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
        "bar": XL_CHART_TYPE.BAR_CLUSTERED,
        "area": XL_CHART_TYPE.AREA,
        "pie": XL_CHART_TYPE.PIE,
    }
    try:
        return values[value]
    except KeyError as exc:
        raise VisualizationRenderError(
            f"MVP chart renderer does not support chart_type={value!r}; use line, column, bar, area, or pie"
        ) from exc


def _validate_chart(visualization: Mapping[str, Any], chart_type: str) -> None:
    categories = visualization.get("categories")
    series = visualization.get("series")
    if not isinstance(categories, list) or not categories:
        raise VisualizationRenderError(
            "chart requires at least one category"
        )
    if not isinstance(series, list) or not series:
        raise VisualizationRenderError(
            "chart requires at least one series"
        )
    if chart_type == "pie" and len(series) != 1:
        raise VisualizationRenderError(
            "pie chart requires exactly one series"
        )

    has_numeric_value = False
    for item in series:
        if not isinstance(item, Mapping):
            raise VisualizationRenderError(
                "each chart series must be an object"
            )
        if item.get("axis", "primary") == "secondary":
            raise VisualizationRenderError(
                "secondary-axis series are not supported by the native renderer"
            )
        values = item.get("values")
        if not isinstance(values, list) or len(values) != len(categories):
            raise VisualizationRenderError(
                "series values must match categories length"
            )
        for value in values:
            if value is None:
                if chart_type == "pie":
                    raise VisualizationRenderError(
                        "pie chart cannot contain missing values"
                    )
                continue
            if (
                isinstance(value, bool)
                or not isinstance(value, (int, float))
                or not math.isfinite(float(value))
            ):
                raise VisualizationRenderError(
                    "chart values must be finite numbers or null"
                )
            if chart_type == "pie" and value < 0:
                raise VisualizationRenderError(
                    "pie chart cannot contain negative values"
                )
            has_numeric_value = True
    if not has_numeric_value:
        raise VisualizationRenderError(
            "chart requires at least one numeric value"
        )
    if chart_type == "pie":
        values = series[0]["values"]
        if sum(float(value) for value in values) <= 0:
            raise VisualizationRenderError(
                "pie chart values must sum to a positive number"
            )

    forecast_start_index = visualization.get("forecast_start_index")
    if (
        forecast_start_index is not None
        and (
            not isinstance(forecast_start_index, int)
            or isinstance(forecast_start_index, bool)
            or not 0 <= forecast_start_index < len(categories)
        )
    ):
        raise VisualizationRenderError(
            "forecast_start_index must reference an existing category"
        )


def _chart_data(visualization: Mapping[str, Any]) -> CategoryChartData:
    categories = visualization.get("categories")
    series = visualization.get("series")
    if not isinstance(categories, list) or not isinstance(series, list):
        raise VisualizationRenderError(
            "chart requires categories and at least one series"
        )
    data = CategoryChartData()
    data.categories = [str(value) for value in categories]
    for item in series:
        if not isinstance(item, Mapping):
            raise VisualizationRenderError("each chart series must be an object")
        values = item.get("values")
        if not isinstance(values, list) or len(values) != len(categories):
            raise VisualizationRenderError("series values must match categories length")
        # The project pins python-pptx 0.6.23, whose chart workbook writer
        # preserves None as a blank cell. Never replace missing evidence with 0.
        data.add_series(str(item.get("name", "Series")), list(values))
    return data


def _style_color(
    style: Mapping[str, Any],
    key: str,
    default: str,
    *,
    context: str = "chart",
) -> RGBColor:
    raw = style.get(key, default)
    value = str(raw).strip().upper()
    if len(value) != 6 or any(
        character not in "0123456789ABCDEF" for character in value
    ):
        raise VisualizationRenderError(
            f"{context} style {key} must be a six-digit RGB hex value"
        )
    return RGBColor.from_string(value)


def _series_colors(style: Mapping[str, Any]) -> tuple[RGBColor, ...]:
    raw = style.get("series_colors", DEFAULT_CHART_COLORS)
    if (
        not isinstance(raw, Sequence)
        or isinstance(raw, (str, bytes))
        or not raw
    ):
        raise VisualizationRenderError(
            "chart style series_colors must be a non-empty array"
        )
    return tuple(
        _style_color({"value": value}, "value", DEFAULT_CHART_COLORS[0])
        for value in raw
    )


def _apply_chart_font(
    font: Any,
    style: Mapping[str, Any],
    *,
    size_key: str,
    default_size: float,
) -> None:
    font.size = Pt(float(style.get(size_key, default_size)))
    font_name = str(style.get("font_family", "")).strip()
    if font_name:
        font.name = font_name
    font.color.rgb = _style_color(
        style,
        "text_color",
        DEFAULT_TEXT_COLOR,
    )


def _style_chart_series(
    chart: Any,
    *,
    chart_type: str,
    visualization: Mapping[str, Any],
    style: Mapping[str, Any],
) -> None:
    palette = _series_colors(style)
    line_width = Pt(float(style.get("line_width_pt", 2.25)))
    marker_size = int(style.get("marker_size_pt", 6))
    for index, series in enumerate(chart.series):
        color = palette[index % len(palette)]
        if chart_type == "line":
            series.format.line.color.rgb = color
            series.format.line.width = line_width
            series.marker.style = XL_MARKER_STYLE.CIRCLE
            series.marker.size = marker_size
            series.marker.format.fill.solid()
            series.marker.format.fill.fore_color.rgb = color
            series.marker.format.line.color.rgb = color
            forecast_start_index = visualization.get("forecast_start_index")
            if isinstance(forecast_start_index, int):
                for point_index in range(
                    forecast_start_index,
                    len(series.points),
                ):
                    point_marker = series.points[point_index].marker
                    point_marker.style = XL_MARKER_STYLE.CIRCLE
                    point_marker.size = marker_size
                    point_marker.format.fill.solid()
                    point_marker.format.fill.fore_color.rgb = RGBColor(
                        255,
                        255,
                        255,
                    )
                    point_marker.format.line.color.rgb = color
        elif chart_type == "pie":
            for point_index, point in enumerate(series.points):
                point_color = palette[point_index % len(palette)]
                point.format.fill.solid()
                point.format.fill.fore_color.rgb = point_color
                point.format.line.color.rgb = RGBColor(255, 255, 255)
        else:
            series.format.fill.solid()
            series.format.fill.fore_color.rgb = color
            series.format.line.color.rgb = color


def _style_chart_axes(
    chart: Any,
    *,
    chart_type: str,
    visualization: Mapping[str, Any],
    style: Mapping[str, Any],
) -> None:
    if chart_type == "pie":
        return
    font_size = float(style.get("font_size_pt", 10))
    axis_color = _style_color(style, "axis_color", DEFAULT_AXIS_COLOR)
    gridline_color = _style_color(
        style,
        "gridline_color",
        DEFAULT_GRIDLINE_COLOR,
    )
    category_axis = chart.category_axis
    value_axis = chart.value_axis
    _apply_chart_font(
        category_axis.tick_labels.font,
        style,
        size_key="font_size_pt",
        default_size=font_size,
    )
    _apply_chart_font(
        value_axis.tick_labels.font,
        style,
        size_key="font_size_pt",
        default_size=font_size,
    )
    category_axis.format.line.color.rgb = axis_color
    value_axis.format.line.color.rgb = axis_color
    value_axis.has_major_gridlines = True
    value_axis.major_gridlines.format.line.color.rgb = gridline_color
    value_axis.major_gridlines.format.line.width = Pt(
        float(style.get("gridline_width_pt", 0.75))
    )

    number_formats = {
        str(item["number_format"])
        for item in visualization.get("series", [])
        if isinstance(item, Mapping) and item.get("number_format")
    }
    if len(number_formats) == 1:
        value_axis.tick_labels.number_format = next(iter(number_formats))
        value_axis.tick_labels.number_format_is_linked = False


def _style_data_labels(
    chart: Any,
    *,
    chart_type: str,
    visualization: Mapping[str, Any],
    style: Mapping[str, Any],
) -> None:
    plot = chart.plots[0]
    if chart_type == "pie":
        show_labels = True
    else:
        show_labels = bool(
            style.get("show_data_labels", chart_type == "bar")
        )
    plot.has_data_labels = show_labels
    if not show_labels:
        return
    labels = plot.data_labels
    _apply_chart_font(
        labels.font,
        style,
        size_key="data_label_font_size_pt",
        default_size=9,
    )
    labels.show_legend_key = False
    labels.show_series_name = False
    labels.show_category_name = False
    if chart_type == "pie":
        labels.show_value = False
        labels.show_percentage = True
        labels.position = XL_LABEL_POSITION.BEST_FIT
        labels.number_format = str(
            style.get("pie_label_number_format", "0%")
        )
        labels.number_format_is_linked = False
        return
    labels.show_value = True
    labels.show_percentage = False
    labels.position = (
        XL_LABEL_POSITION.OUTSIDE_END
        if chart_type in {"column", "bar"}
        else XL_LABEL_POSITION.ABOVE
    )
    number_formats = {
        str(item["number_format"])
        for item in visualization.get("series", [])
        if isinstance(item, Mapping) and item.get("number_format")
    }
    if len(number_formats) == 1:
        labels.number_format = next(iter(number_formats))
        labels.number_format_is_linked = False


def _style_chart_legend(
    chart: Any,
    *,
    chart_type: str,
    series_count: int,
    style: Mapping[str, Any],
) -> None:
    chart.has_legend = chart_type == "pie" or series_count > 1
    if not chart.has_legend:
        return
    positions = {
        "bottom": XL_LEGEND_POSITION.BOTTOM,
        "left": XL_LEGEND_POSITION.LEFT,
        "right": XL_LEGEND_POSITION.RIGHT,
        "top": XL_LEGEND_POSITION.TOP,
    }
    position_name = str(style.get("legend_position", "bottom")).strip().lower()
    if position_name not in positions:
        raise VisualizationRenderError(
            "chart style legend_position must be bottom, left, right, or top"
        )
    chart.legend.position = positions[position_name]
    chart.legend.include_in_layout = False
    _apply_chart_font(
        chart.legend.font,
        style,
        size_key="legend_font_size_pt",
        default_size=9,
    )


def _normalize_chart_axis_ids(chart: Any) -> None:
    """Keep generated axis IDs within the unsigned range required by Open XML."""

    for axis_id in chart._chartSpace.xpath(".//c:axId | .//c:crossAx"):
        raw = axis_id.get("val")
        if raw is None:
            continue
        value = int(raw)
        if value < 0 or value > 0x7FFFFFFF:
            axis_id.set("val", str(value & 0x7FFFFFFF))


def _set_display_blanks_as(chart: Any, value: str) -> None:
    """Set the native chart empty-cell display mode without changing data."""

    chart_element = chart._chartSpace.chart
    display_blanks = chart_element.find(qn("c:dispBlanksAs"))
    if display_blanks is None:
        display_blanks = OxmlElement("c:dispBlanksAs")
        successor = next(
            (
                chart_element.find(qn(tag))
                for tag in ("c:showDLblsOverMax", "c:extLst")
                if chart_element.find(qn(tag)) is not None
            ),
            None,
        )
        if successor is None:
            chart_element.append(display_blanks)
        else:
            successor.addprevious(display_blanks)
    display_blanks.set("val", value)


def remove_overlapping_charts(slide: Any, anchor: Any) -> None:
    """Remove template example charts whose bounds overlap a semantic anchor."""

    ax, ay, aw, ah = anchor.left, anchor.top, anchor.width, anchor.height
    for shape in list(slide.shapes):
        if not getattr(shape, "has_chart", False):
            continue
        sx, sy, sw, sh = shape.left, shape.top, shape.width, shape.height
        overlap = max(0, min(ax + aw, sx + sw) - max(ax, sx)) * max(
            0, min(ay + ah, sy + sh) - max(ay, sy)
        )
        if overlap > 0:
            slide.shapes._spTree.remove(shape._element)


def render_chart(
    slide: Any,
    anchor: Any,
    visualization: Mapping[str, Any],
    *,
    style: Mapping[str, Any] | None = None,
) -> Any:
    remove_overlapping_charts(slide, anchor)
    chart_type = str(visualization.get("chart_type"))
    _chart_type(chart_type)
    _validate_chart(visualization, chart_type)
    chart = slide.shapes.add_chart(
        _chart_type(chart_type),
        anchor.left,
        anchor.top,
        anchor.width,
        anchor.height,
        _chart_data(visualization),
    ).chart
    _normalize_chart_axis_ids(chart)
    if chart_type == "line":
        # Keep the embedded cell blank and omit its marker, but span the line
        # between adjacent known points instead of showing a broken segment.
        _set_display_blanks_as(chart, "span")
    chart.has_title = False
    style = style or {}
    plot = chart.plots[0]
    if chart_type in {"column", "bar"}:
        plot.gap_width = int(style.get("gap_width", 65))
    _style_chart_series(
        chart,
        chart_type=chart_type,
        visualization=visualization,
        style=style,
    )
    _style_chart_axes(
        chart,
        chart_type=chart_type,
        visualization=visualization,
        style=style,
    )
    _style_data_labels(
        chart,
        chart_type=chart_type,
        visualization=visualization,
        style=style,
    )
    _style_chart_legend(
        chart,
        chart_type=chart_type,
        series_count=len(visualization.get("series", [])),
        style=style,
    )
    return chart


def _display_width(value: object) -> int:
    text = "" if value is None else str(value)
    return sum(
        2
        if unicodedata.east_asian_width(character) in {"W", "F", "A"}
        else 1
        for character in text
    )


_NUMERIC_LIKE = re.compile(
    r"^[+-]?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?(?:%|倍|x|X)?$"
)


def _is_numeric_like(value: object) -> bool:
    if isinstance(value, bool) or value is None:
        return False
    if isinstance(value, (int, float)):
        return math.isfinite(float(value))
    return bool(_NUMERIC_LIKE.fullmatch(str(value).strip()))


def _table_column_widths(
    total_width: int,
    columns: Sequence[object],
    rows: Sequence[Sequence[object]],
) -> tuple[int, ...]:
    column_count = len(columns)
    weights: list[float] = []
    for column_index, heading in enumerate(columns):
        values = [
            row[column_index]
            for row in rows
            if row[column_index] is not None
            and str(row[column_index]).strip()
        ]
        numeric = bool(values) and all(
            _is_numeric_like(value) for value in values
        )
        content_width = max(
            [
                _display_width(heading),
                *[_display_width(value) for value in values],
            ]
        )
        if numeric:
            weight = float(max(8, min(content_width + 2, 14)))
        else:
            weight = float(max(12, min(content_width + 2, 32)))
        if column_index == 0 and not numeric:
            weight *= 1.25
        weights.append(weight)

    minimum_share = 0.06 if column_count >= 6 else 0.08
    minimum_width = int(total_width * minimum_share)
    distributable = max(
        0,
        total_width - minimum_width * column_count,
    )
    total_weight = sum(weights)
    widths = [
        minimum_width + int(distributable * weight / total_weight)
        for weight in weights
    ]
    widths[-1] += total_width - sum(widths)
    return tuple(widths)


def _set_table_cell_border(
    cell: Any,
    *,
    color: RGBColor,
    width_pt: float,
) -> None:
    properties = cell._tc.get_or_add_tcPr()
    for edge in ("lnL", "lnR", "lnT", "lnB"):
        tag = qn(f"a:{edge}")
        existing = properties.find(tag)
        if existing is not None:
            properties.remove(existing)
        line = OxmlElement(f"a:{edge}")
        line.set("w", str(int(Pt(width_pt))))
        solid_fill = OxmlElement("a:solidFill")
        rgb = OxmlElement("a:srgbClr")
        rgb.set("val", str(color))
        solid_fill.append(rgb)
        line.append(solid_fill)
        dash = OxmlElement("a:prstDash")
        dash.set("val", "solid")
        line.append(dash)
        successor = next(
            (
                properties.find(qn(f"a:{name}"))
                for name in (
                    "solidFill",
                    "gradFill",
                    "blipFill",
                    "pattFill",
                    "grpFill",
                    "headers",
                    "extLst",
                )
                if properties.find(qn(f"a:{name}")) is not None
            ),
            None,
        )
        if successor is None:
            properties.append(line)
        else:
            successor.addprevious(line)


def _style_table_cell(
    cell: Any,
    *,
    value: object,
    header: bool,
    first_column: bool,
    striped: bool,
    style: Mapping[str, Any],
) -> None:
    header_fill = _style_color(
        style,
        "header_fill",
        DEFAULT_TABLE_HEADER_FILL,
        context="table",
    )
    header_text = _style_color(
        style,
        "header_text_color",
        DEFAULT_TABLE_HEADER_TEXT,
        context="table",
    )
    body_fill = _style_color(
        style,
        "body_fill",
        DEFAULT_TABLE_BODY_FILL,
        context="table",
    )
    stripe_fill = _style_color(
        style,
        "stripe_fill",
        DEFAULT_TABLE_STRIPE_FILL,
        context="table",
    )
    body_text = _style_color(
        style,
        "text_color",
        DEFAULT_TABLE_TEXT,
        context="table",
    )
    border_color = _style_color(
        style,
        "border_color",
        DEFAULT_TABLE_BORDER,
        context="table",
    )

    cell.fill.solid()
    cell.fill.fore_color.rgb = (
        header_fill
        if header
        else stripe_fill
        if striped
        else body_fill
    )
    cell.margin_left = Pt(float(style.get("margin_horizontal_pt", 5)))
    cell.margin_right = Pt(float(style.get("margin_horizontal_pt", 5)))
    cell.margin_top = Pt(float(style.get("margin_vertical_pt", 3)))
    cell.margin_bottom = Pt(float(style.get("margin_vertical_pt", 3)))
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.text_frame.word_wrap = True
    alignment = PP_ALIGN.CENTER
    font_name = str(style.get("font_family", "")).strip() or None
    font_size = float(style.get("font_size_pt", 9))
    for paragraph in cell.text_frame.paragraphs:
        paragraph.alignment = alignment
        for run in paragraph.runs:
            run.font.size = Pt(font_size)
            run.font.bold = bool(
                header
                or (
                    first_column
                    and style.get("first_column_bold", True)
                )
            )
            if font_name:
                run.font.name = font_name
            run.font.color.rgb = header_text if header else body_text
    _set_table_cell_border(
        cell,
        color=border_color,
        width_pt=float(style.get("border_width_pt", 0.6)),
    )


def render_table(
    slide: Any,
    target: Any,
    visualization: Mapping[str, Any],
    *,
    style: Mapping[str, Any] | None = None,
) -> Any:
    columns = visualization.get("columns")
    rows = visualization.get("rows")
    if (
        not isinstance(columns, list)
        or not isinstance(rows, list)
        or not columns
        or not rows
    ):
        raise VisualizationRenderError("table requires columns and rows")
    if any(not isinstance(row, list) or len(row) != len(columns) for row in rows):
        raise VisualizationRenderError("each table row must match columns length")
    style = style or {}
    max_rows = int(style.get("max_rows", DEFAULT_TABLE_MAX_ROWS))
    max_columns = int(style.get("max_columns", DEFAULT_TABLE_MAX_COLUMNS))
    if len(rows) > max_rows:
        raise VisualizationRenderError(
            f"table exceeds renderer max_rows={max_rows}; "
            "split the table or use a continuation page"
        )
    if len(columns) > max_columns:
        raise VisualizationRenderError(
            f"table exceeds renderer max_columns={max_columns}; "
            "use fewer columns or a continuation page"
        )
    left, top, width, height = target.left, target.top, target.width, target.height
    old_element = getattr(target, "_element", None)
    target_name = getattr(target, "name", "compiled_table")
    if old_element is not None:
        slide.shapes._spTree.remove(old_element)
    desired_header_height = int(
        Pt(float(style.get("header_height_pt", 30)))
    )
    desired_row_height = int(Pt(float(style.get("row_height_pt", 26))))
    desired_height = desired_header_height + desired_row_height * len(rows)
    rendered_height = min(height, desired_height)
    scale = min(1.0, rendered_height / desired_height)
    header_height = int(desired_header_height * scale)
    row_height = int(desired_row_height * scale)
    rendered_height = header_height + row_height * len(rows)
    table_shape = slide.shapes.add_table(
        len(rows) + 1,
        len(columns),
        left,
        top,
        width,
        rendered_height,
    )
    table_shape.name = target_name
    table = table_shape.table
    for column, column_width in zip(
        table.columns,
        _table_column_widths(width, columns, rows),
    ):
        column.width = column_width
    table.rows[0].height = header_height
    for row_index in range(1, len(table.rows)):
        table.rows[row_index].height = row_height
    for col, value in enumerate(columns):
        table.cell(0, col).text = str(value)
        _style_table_cell(
            table.cell(0, col),
            value=value,
            header=True,
            first_column=col == 0,
            striped=False,
            style=style,
        )
    for row_index, row in enumerate(rows, start=1):
        for col, value in enumerate(row):
            table.cell(row_index, col).text = "" if value is None else str(value)
            _style_table_cell(
                table.cell(row_index, col),
                value=value,
                header=False,
                first_column=col == 0,
                striped=row_index % 2 == 0,
                style=style,
            )
    return table_shape


def resolve_image_path(asset_path: object, *, asset_root: Path | None = None) -> Path:
    """Resolve a Visualization image within an explicit/current asset root."""

    if not isinstance(asset_path, str) or not asset_path.strip():
        raise VisualizationRenderError("image requires a non-empty asset_path")
    relative = Path(asset_path)
    if relative.is_absolute() or ".." in relative.parts:
        raise VisualizationRenderError("image asset_path must be a safe relative path")
    root = (asset_root or Path.cwd()).resolve()
    resolved = (root / relative).resolve()
    try:
        resolved.relative_to(root)
    except ValueError as exc:
        raise VisualizationRenderError("image asset_path escapes the asset root") from exc
    if not resolved.is_file():
        raise VisualizationRenderError(f"image asset does not exist: {asset_path}")
    return resolved


def render_image(
    slide: Any,
    anchor: Any,
    visualization: Mapping[str, Any],
    *,
    asset_root: Path | None = None,
) -> Any:
    """Insert a source image into a Layout Map anchor without distortion."""

    path = resolve_image_path(visualization.get("asset_path"), asset_root=asset_root)
    try:
        with PILImage.open(path) as image:
            pixel_width, pixel_height = image.size
            image.verify()
    except (OSError, ValueError) as exc:
        raise VisualizationRenderError(f"image asset is invalid: {path.name}") from exc
    if pixel_width <= 0 or pixel_height <= 0:
        raise VisualizationRenderError(f"image asset has invalid dimensions: {path.name}")

    scale = min(anchor.width / pixel_width, anchor.height / pixel_height)
    width = max(1, int(pixel_width * scale))
    height = max(1, int(pixel_height * scale))
    left = anchor.left + (anchor.width - width) // 2
    top = anchor.top + (anchor.height - height) // 2
    picture = slide.shapes.add_picture(str(path), left, top, width, height)
    source = visualization.get("source")
    identity = source.get("id") if isinstance(source, Mapping) else None
    picture.name = f"image_{identity or path.stem}"
    return picture

"""Template slide cloning and semantic field population."""

from __future__ import annotations

from copy import deepcopy
from pathlib import Path
from types import SimpleNamespace
from typing import Any, Iterable, Mapping, Sequence

from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN

from .visualization_renderer import render_chart, render_image, render_table


class SlideBuildError(ValueError):
    """Raised when a template field cannot be populated."""


def duplicate_slide(prs: Any, source_slide: Any) -> Any:
    """Append a template slide without retaining native example data."""

    target = prs.slides.add_slide(source_slide.slide_layout)
    sp_tree = target.shapes._spTree
    for shape in list(target.shapes):
        sp_tree.remove(shape._element)
    for shape in source_slide.shapes:
        if getattr(shape, "has_chart", False):
            continue
        if getattr(shape, "has_table", False):
            anchor = target.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                shape.left,
                shape.top,
                shape.width,
                shape.height,
            )
            anchor.name = shape.name
            anchor.fill.background()
            anchor.line.fill.background()
            continue
        sp_tree.append(deepcopy(shape.element))
    source_bg = source_slide.element.cSld.find(qn("p:bg"))
    if source_bg is not None:
        target_bg = target.element.cSld.find(qn("p:bg"))
        if target_bg is not None:
            target.element.cSld.remove(target_bg)
        target.element.cSld.insert(0, deepcopy(source_bg))
    return target


def remove_slide(prs: Any, index: int) -> None:
    slide_id = prs.slides._sldIdLst[index]
    prs.part.drop_rel(slide_id.rId)
    prs.slides._sldIdLst.remove(slide_id)


def _target_name(spec: Mapping[str, Any]) -> str | None:
    target = spec.get("target")
    return target.get("name") if isinstance(target, Mapping) else target if isinstance(target, str) else None


def find_shape(slide: Any, target: Mapping[str, Any] | str | None) -> Any | None:
    if target is None:
        return None
    if isinstance(target, Mapping):
        name = target.get("name")
        shape_id = target.get("shape_id")
    else:
        name = target
        shape_id = None
    if isinstance(name, str):
        for shape in slide.shapes:
            if shape.name == name:
                return shape
    if shape_id is not None:
        for shape in slide.shapes:
            if shape.shape_id == shape_id:
                return shape
    return None


def set_text(shape: Any | None, value: Any, *, required: bool = False) -> None:
    if shape is None:
        if required:
            raise SlideBuildError("required template text target is missing")
        return
    shape.text = "" if value is None else str(value)


def set_bullets(shape: Any | None, values: Iterable[Any], *, required: bool = False) -> None:
    if shape is None:
        if required:
            raise SlideBuildError("required template bullet target is missing")
        return
    items = [str(value) for value in values if value is not None and str(value).strip()]
    shape.text = "\n".join(items)


def _remove_named_shapes(slide: Any, names: Sequence[str]) -> None:
    targets = set(names)
    for shape in list(slide.shapes):
        if shape.name in targets:
            shape._element.getparent().remove(shape._element)


def _compiled_target(slide: Any, target: Mapping[str, Any]) -> Any:
    bounds = target.get("bounds_in")
    if isinstance(bounds, Mapping):
        return SimpleNamespace(
            left=Inches(float(bounds["left"])),
            top=Inches(float(bounds["top"])),
            width=Inches(float(bounds["width"])),
            height=Inches(float(bounds["height"])),
        )
    shape = find_shape(slide, target)
    if shape is None:
        raise SlideBuildError(f"compiled target is missing: {dict(target)!r}")
    return shape


def _apply_compiled_text_style(shape: Any, style: Mapping[str, Any]) -> None:
    alignment = {
        "left": PP_ALIGN.LEFT,
        "center": PP_ALIGN.CENTER,
        "right": PP_ALIGN.RIGHT,
    }
    vertical = {
        "top": MSO_ANCHOR.TOP,
        "middle": MSO_ANCHOR.MIDDLE,
        "bottom": MSO_ANCHOR.BOTTOM,
    }
    frame = shape.text_frame
    frame.vertical_anchor = vertical.get(
        str(style.get("vertical_alignment", "top")), MSO_ANCHOR.TOP
    )
    for paragraph in frame.paragraphs:
        paragraph.alignment = alignment.get(
            str(style.get("alignment", "left")), PP_ALIGN.LEFT
        )
        for run in paragraph.runs:
            run.font.name = str(style["font_family"])
            run.font.size = Pt(float(style["font_size_pt"]))
            run.font.bold = bool(style.get("bold", False))
            run.font.color.rgb = RGBColor.from_string(str(style["color"]))


def _add_compiled_text_box(
    slide: Any,
    target: Mapping[str, Any],
    value: Any,
    style: Mapping[str, Any],
    *,
    bullets: bool = False,
) -> Any:
    anchor = _compiled_target(slide, target)
    shape = slide.shapes.add_textbox(
        anchor.left, anchor.top, anchor.width, anchor.height
    )
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    frame.margin_left = Pt(2)
    frame.margin_right = Pt(2)
    frame.margin_top = Pt(1)
    frame.margin_bottom = Pt(1)
    values = value if bullets and isinstance(value, Sequence) else [value]
    for index, item in enumerate(values):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = str(item)
        paragraph.level = 0
        if bullets:
            paragraph.text = f"•{item}"
    _apply_compiled_text_style(shape, style)
    return shape


def _set_compiled_repeated(
    slide: Any,
    targets: Sequence[Mapping[str, Any]],
    values: Sequence[Any],
) -> None:
    for index, target_group in enumerate(targets):
        value = values[index] if index < len(values) else None
        if isinstance(value, Mapping):
            for field, target in target_group.items():
                shape = _compiled_target(slide, target)
                set_text(shape, value.get(field, ""), required=True)
        else:
            fields = list(target_group)
            for field_index, field in enumerate(fields):
                shape = _compiled_target(slide, target_group[field])
                text = value if field_index == len(fields) - 1 and value is not None else ""
                set_text(shape, text, required=True)


def execute_compiled_operations(
    slide: Any,
    operations: Sequence[Mapping[str, Any]],
    *,
    visualizations_by_id: Mapping[str, Mapping[str, Any]],
    asset_root: Path,
) -> None:
    """Execute a complete operation stream without resolving layout semantics."""

    for operation in operations:
        op = str(operation.get("op", ""))
        if op == "add_text_box":
            _add_compiled_text_box(
                slide,
                operation["target"],
                operation.get("value"),
                operation["style"],
            )
        elif op == "add_bullet_list":
            _add_compiled_text_box(
                slide,
                operation["target"],
                operation.get("values", []),
                operation["style"],
                bullets=True,
            )
        elif op == "set_text":
            set_text(
                _compiled_target(slide, operation["target"]),
                operation.get("value"),
                required=True,
            )
        elif op == "set_bullets":
            set_bullets(
                _compiled_target(slide, operation["target"]),
                operation.get("values", []),
                required=True,
            )
        elif op == "set_repeated_text":
            _set_compiled_repeated(
                slide,
                operation.get("targets", []),
                operation.get("values", []),
            )
        elif op == "remove_shapes":
            _remove_named_shapes(slide, operation.get("names", []))
        elif op in {"render_chart", "render_table", "render_image"}:
            visualization_id = str(operation.get("visualization_id", ""))
            record = visualizations_by_id.get(visualization_id)
            if record is None:
                raise SlideBuildError(
                    f"compiled operation references missing visualization {visualization_id!r}"
                )
            data = record.get("data")
            if not isinstance(data, Mapping):
                raise SlideBuildError(
                    f"compiled visualization {visualization_id!r} has no data"
                )
            target = _compiled_target(slide, operation["target"])
            if op == "render_chart":
                render_chart(slide, target, data, style=operation.get("style"))
                title_target = operation.get("title_target")
                if isinstance(title_target, Mapping):
                    set_text(
                        _compiled_target(slide, title_target),
                        data.get("title", ""),
                        required=True,
                    )
            elif op == "render_table":
                render_table(slide, target, data, style=operation.get("style"))
            else:
                render_image(slide, target, data, asset_root=asset_root)
        else:
            raise SlideBuildError(f"unsupported compiled operation: {op!r}")


def _field(layout: Mapping[str, Any], name: str) -> Mapping[str, Any]:
    fields = layout.get("fields", {})
    value = fields.get(name, {}) if isinstance(fields, Mapping) else {}
    return value if isinstance(value, Mapping) else {}


def _field_shape(slide: Any, layout: Mapping[str, Any], name: str) -> Any | None:
    return find_shape(slide, _field(layout, name).get("target"))


def _repeated_targets(layout: Mapping[str, Any], field_name: str) -> Sequence[Mapping[str, Any]]:
    targets = _field(layout, field_name).get("targets", [])
    return targets if isinstance(targets, list) else []


def _set_standard(slide: Any, layout: Mapping[str, Any], outline_slide: Mapping[str, Any], page_number: int) -> None:
    set_text(_field_shape(slide, layout, "title"), outline_slide.get("title"))
    set_text(_field_shape(slide, layout, "page_number"), page_number)
    refs = outline_slide.get("source_refs", [])
    set_text(_field_shape(slide, layout, "source_refs"), ", ".join(str(ref) for ref in refs))


def _set_repeated_text(slide: Any, targets: Sequence[Mapping[str, Any]], items: Sequence[Any], fields: Sequence[str]) -> None:
    for index, target_group in enumerate(targets):
        value = items[index] if index < len(items) else None
        if isinstance(value, Mapping):
            for field in fields:
                target = target_group.get(field)
                set_text(find_shape(slide, target), value.get(field, ""))
        else:
            target = target_group.get(fields[-1])
            set_text(find_shape(slide, target), value if index < len(items) else "")


def _image_anchors(slide: Any, layout: Mapping[str, Any]) -> list[Any]:
    """Return Layout Map anchors in image-specific then compatible order."""

    result: list[Any] = []
    seen: set[int] = set()
    image_field = _field(layout, "image")
    bounds = image_field.get("bounds_in")
    if isinstance(bounds, Mapping) and all(
        isinstance(bounds.get(key), (int, float))
        for key in ("left", "top", "width", "height")
    ):
        result.append(
            SimpleNamespace(
                left=Inches(float(bounds["left"])),
                top=Inches(float(bounds["top"])),
                width=Inches(float(bounds["width"])),
                height=Inches(float(bounds["height"])),
            )
        )
        return result
    for name in ("image", "chart", "chart_left", "chart_right", "table"):
        field = _field(layout, name)
        target = field.get("anchor") or field.get("target")
        shape = find_shape(slide, target)
        if shape is not None and shape.shape_id not in seen:
            result.append(shape)
            seen.add(shape.shape_id)
    if result:
        return result

    # Compatibility fallback for layouts designed before image slots existed.
    # Prefer the union of repeated content targets so an image can occupy the
    # layout's main body rather than a narrow headline box. Coordinates still
    # come entirely from the Layout Map/template, never renderer constants.
    excluded = {"title", "page_number", "source_refs", "cover_title", "cover_meta", "tag"}
    fields = layout.get("fields", {})
    if isinstance(fields, Mapping):
        repeated_shapes = []
        candidates = []
        for name, value in fields.items():
            if name in excluded or not isinstance(value, Mapping):
                continue
            shape = find_shape(slide, value.get("target"))
            if shape is not None:
                candidates.append(shape)
            targets = value.get("targets")
            if isinstance(targets, list):
                for group in targets:
                    if not isinstance(group, Mapping):
                        continue
                    for target in group.values():
                        shape = find_shape(slide, target)
                        if shape is not None:
                            repeated_shapes.append(shape)
        if repeated_shapes:
            left = min(shape.left for shape in repeated_shapes)
            top = min(shape.top for shape in repeated_shapes)
            right = max(shape.left + shape.width for shape in repeated_shapes)
            bottom = max(shape.top + shape.height for shape in repeated_shapes)
            result.append(
                SimpleNamespace(
                    left=left,
                    top=top,
                    width=right - left,
                    height=bottom - top,
                )
            )
            return result
        if candidates:
            result.append(max(candidates, key=lambda shape: shape.width * shape.height))
    return result


def _render_images(
    slide: Any,
    layout: Mapping[str, Any],
    image_visuals: Sequence[Mapping[str, Any]],
    *,
    occupied_slots: int = 0,
    asset_root: Path | None = None,
) -> None:
    anchors = _image_anchors(slide, layout)
    for index, visualization in enumerate(image_visuals, start=occupied_slots):
        if index >= len(anchors):
            raise SlideBuildError(
                f"layout {layout.get('layout_id', '')!r} has no image-compatible slot for visualization {index + 1}"
            )
        render_image(slide, anchors[index], visualization, asset_root=asset_root)


def populate_slide(
    slide: Any,
    layout: Mapping[str, Any],
    outline_slide: Mapping[str, Any],
    *,
    metadata: Mapping[str, Any],
    visualizations: Sequence[Mapping[str, Any]],
    page_number: int,
    asset_root: Path | None = None,
) -> None:
    """Populate the MVP fields for a resolved semantic layout."""

    layout_id = str(layout.get("layout_id", ""))
    _set_standard(slide, layout, outline_slide, page_number)
    title = outline_slide.get("title", "")
    key_message = outline_slide.get("key_message", "")
    bullets = outline_slide.get("bullet_points", [])
    image_visuals = [item for item in visualizations if item.get("type") == "image"]
    if not isinstance(bullets, list):
        bullets = []

    if layout_id == "image_only_layout":
        if len(image_visuals) != 1:
            raise SlideBuildError(
                "image_only_layout requires exactly one original figure image"
            )
        set_text(find_shape(slide, "layout_type"), "ORIGINAL FIGURE")
        _remove_named_shapes(
            slide,
            (
                "logic_chart_panel",
                "logic_chart_title",
                "logic_takeaway",
                "logic_main_point",
                "logic_bullets",
                "logic_conclusion",
                "logic_conclusion_text",
            ),
        )
        _render_images(
            slide,
            layout,
            image_visuals,
            asset_root=asset_root,
        )
        return

    if layout_id == "cover":
        set_text(_field_shape(slide, layout, "cover_title"), metadata.get("company_name", ""))
        set_text(_field_shape(slide, layout, "subtitle"), metadata.get("report_title", title))
        meta = "  |  ".join(str(metadata.get(key, "")) for key in ("stock_code", "report_date") if metadata.get(key))
        set_text(_field_shape(slide, layout, "cover_meta"), meta)
        set_text(_field_shape(slide, layout, "tag"), metadata.get("industry", ""))
        _render_images(slide, layout, image_visuals, asset_root=asset_root)
        return

    if layout_id == "executive_summary":
        set_text(_field_shape(slide, layout, "thesis"), key_message)
        _set_repeated_text(slide, _repeated_targets(layout, "logics"), bullets, ("title", "body"))
        _render_images(slide, layout, image_visuals, asset_root=asset_root)
        return
    if layout_id == "company_overview":
        set_text(_field_shape(slide, layout, "company_name"), metadata.get("company_name"))
        set_text(_field_shape(slide, layout, "company_code"), metadata.get("stock_code"))
        set_text(_field_shape(slide, layout, "company_positioning"), key_message)
        _set_repeated_text(slide, _repeated_targets(layout, "segments"), bullets, ("title", "body"))
    elif layout_id in {"industry_outlook", "competitive_landscape", "risk_catalyst", "capability_map"}:
        field = {"industry_outlook": "drivers", "competitive_landscape": "moats", "risk_catalyst": "risks", "capability_map": "stages"}[layout_id]
        _set_repeated_text(slide, _repeated_targets(layout, field), bullets, ("title", "body"))
    elif layout_id == "chart_text":
        set_text(_field_shape(slide, layout, "takeaway"), key_message)
        set_bullets(_field_shape(slide, layout, "bullets"), bullets)
        set_text(_field_shape(slide, layout, "conclusion"), key_message)
    elif layout_id == "valuation":
        set_text(_field_shape(slide, layout, "headline"), key_message)
        set_bullets(_field_shape(slide, layout, "bullets"), bullets)
        set_text(_field_shape(slide, layout, "conclusion"), key_message)
    elif layout_id == "earnings_forecast":
        set_text(_field_shape(slide, layout, "note"), key_message)
        set_bullets(_field_shape(slide, layout, "assumptions"), bullets)

    chart_visuals = [item for item in visualizations if "chart_type" in item]
    table_visuals = [item for item in visualizations if "columns" in item]
    chart_fields = [
        name for name in ("chart", "chart_left", "chart_right")
        if isinstance(_field(layout, name).get("anchor"), Mapping)
        and find_shape(slide, _field(layout, name).get("anchor")) is not None
    ]
    for index, visualization in enumerate(chart_visuals):
        if index >= len(chart_fields):
            raise SlideBuildError(f"layout {layout_id!r} has no chart slot for visualization {index + 1}")
        field = _field(layout, chart_fields[index])
        anchor = find_shape(slide, field.get("anchor"))
        if anchor is None:
            raise SlideBuildError(f"layout {layout_id!r} chart anchor is missing")
        render_chart(slide, anchor, visualization)
        set_text(find_shape(slide, field.get("title_target")), visualization.get("title", ""))
    for visualization in table_visuals:
        field = _field(layout, "table")
        target = find_shape(slide, field.get("target"))
        if target is None:
            raise SlideBuildError(f"layout {layout_id!r} table target is missing")
        render_table(slide, target, visualization)
    _render_images(
        slide,
        layout,
        image_visuals,
        occupied_slots=len(chart_visuals) + len(table_visuals),
        asset_root=asset_root,
    )

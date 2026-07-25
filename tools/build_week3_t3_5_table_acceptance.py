"""Build the fixed four-slide native-table acceptance deck for T3.5."""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from types import SimpleNamespace
from typing import Any, Mapping, Sequence

from jsonschema import Draft202012Validator
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from ppt_engine.visualization_renderer import render_table


EXAMPLE_ROOT = PROJECT_ROOT / "examples/visualization_tables"
DEFAULT_EXAMPLES = (
    EXAMPLE_ROOT / "week3_t3_5_forecast.json",
    EXAMPLE_ROOT / "week3_t3_5_valuation.json",
    EXAMPLE_ROOT / "week3_t3_5_long_labels.json",
    EXAMPLE_ROOT / "week3_t3_5_capacity_boundary.json",
)
DEFAULT_OUTPUT = PROJECT_ROOT / "output/week3_t3.5_table_acceptance.pptx"
VISUALIZATION_SCHEMA = PROJECT_ROOT / "schemas/visualization.schema.json"
ACCEPTANCE_STYLE = {
    "font_family": "Microsoft YaHei",
    "font_size_pt": 9,
    "text_color": "334155",
    "header_fill": "102A43",
    "header_text_color": "FFFFFF",
    "body_fill": "FFFFFF",
    "stripe_fill": "F5F7FA",
    "border_color": "D9E2EC",
    "border_width_pt": 0.6,
    "margin_horizontal_pt": 5,
    "margin_vertical_pt": 3,
    "header_height_pt": 30,
    "row_height_pt": 26,
    "first_column_bold": True,
    "max_rows": 18,
    "max_columns": 8,
}


def _load_table(path: Path, validator: Draft202012Validator) -> dict[str, Any]:
    value = json.loads(path.read_text(encoding="utf-8"))
    errors = sorted(
        validator.iter_errors(value),
        key=lambda error: list(error.absolute_path),
    )
    if errors:
        raise ValueError(f"{path.name}: {errors[0].message}")
    if "columns" not in value or "rows" not in value:
        raise ValueError(f"{path.name}: example is not a table")
    return value


def _add_text(
    slide: Any,
    text: str,
    *,
    left: float,
    top: float,
    width: float,
    height: float,
    size: float,
    color: str,
    bold: bool = False,
    align: Any = PP_ALIGN.LEFT,
) -> None:
    shape = slide.shapes.add_textbox(
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    shape.text_frame.clear()
    paragraph = shape.text_frame.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = align
    run = paragraph.runs[0]
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def build_acceptance_deck(
    tables: Sequence[Mapping[str, Any]],
    output: Path,
) -> None:
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    blank_layout = presentation.slide_layouts[6]

    for index, table_data in enumerate(tables, start=1):
        slide = presentation.slides.add_slide(blank_layout)
        _add_text(
            slide,
            str(table_data["title"]),
            left=0.72,
            top=0.35,
            width=11.9,
            height=0.5,
            size=24,
            color="102A43",
            bold=True,
        )
        _add_text(
            slide,
            f"单位：{table_data.get('unit') or '—'}",
            left=0.76,
            top=0.94,
            width=3.2,
            height=0.3,
            size=10,
            color="64748B",
        )
        anchor = SimpleNamespace(
            left=Inches(0.75),
            top=Inches(1.25),
            width=Inches(11.83),
            height=Inches(5.45),
        )
        render_table(
            slide,
            anchor,
            table_data,
            style=ACCEPTANCE_STYLE,
        )
        _add_text(
            slide,
            (
                f"T3.5 样式验收 {index}/4 · "
                f"{len(table_data['rows'])} 行 × "
                f"{len(table_data['columns'])} 列 · PowerPoint 原生表格"
            ),
            left=0.75,
            top=6.9,
            width=11.83,
            height=0.25,
            size=8.5,
            color="94A3B8",
            align=PP_ALIGN.RIGHT,
        )

    output.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(output)


def parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Build the Week 3 T3.5 native-table acceptance deck"
    )
    parser.add_argument(
        "--example",
        action="append",
        type=Path,
        dest="examples",
    )
    parser.add_argument("-o", "--output", type=Path, default=DEFAULT_OUTPUT)
    return parser.parse_args(argv)


def main(argv: Sequence[str] | None = None) -> int:
    args = parse_args(argv)
    try:
        schema = json.loads(
            VISUALIZATION_SCHEMA.read_text(encoding="utf-8")
        )
        validator = Draft202012Validator(schema)
        paths = tuple(args.examples or DEFAULT_EXAMPLES)
        tables = [_load_table(path, validator) for path in paths]
        build_acceptance_deck(tables, args.output)
        print(f"Created T3.5 acceptance deck: {args.output}")
        print(f"Native table samples: {len(tables)}")
        return 0
    except (OSError, ValueError, json.JSONDecodeError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())

#!/usr/bin/env python3
"""Build a deterministic Template Profile from the existing Layout Map.

The legacy Layout Map remains unchanged during migration.  This builder makes
the renderer capabilities that are actually supported today explicit and
validates the result against ``template_profile.schema.json``.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import sys
from pathlib import Path
from typing import Any, Mapping, Sequence

from jsonschema import Draft202012Validator
from pptx import Presentation
from pptx.enum.dml import MSO_COLOR_TYPE


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SCHEMA = PROJECT_ROOT / "schemas" / "template_profile.schema.json"

SUPPORTED_OPERATIONS = (
    "set_text",
    "set_bullets",
    "set_repeated_text",
    "render_chart",
    "render_table",
    "render_image",
    "remove_shapes",
    "add_text_box",
    "add_bullet_list",
    "add_shape",
    "add_line",
)

_SOURCE_BINDINGS: dict[tuple[str, str], tuple[str, str]] = {
    ("*", "title"): ("slide.title", "set_text"),
    ("*", "page_number"): ("compiler.page_number", "set_text"),
    ("*", "source_refs"): ("slide.source_refs", "set_text"),
    ("cover", "cover_title"): ("metadata.company_name", "set_text"),
    ("cover", "subtitle"): ("metadata.report_title", "set_text"),
    ("cover", "cover_meta"): ("compiler.cover_meta", "set_text"),
    ("cover", "tag"): ("metadata.industry", "set_text"),
    ("executive_summary", "thesis"): ("slide.key_message", "set_text"),
    ("executive_summary", "logics"): ("slide.bullet_points", "set_repeated_text"),
    ("company_overview", "company_name"): ("metadata.company_name", "set_text"),
    ("company_overview", "company_code"): ("metadata.stock_code", "set_text"),
    ("company_overview", "company_positioning"): ("slide.key_message", "set_text"),
    ("company_overview", "segments"): ("slide.bullet_points", "set_repeated_text"),
    ("industry_outlook", "drivers"): ("slide.bullet_points", "set_repeated_text"),
    ("competitive_landscape", "moats"): ("slide.bullet_points", "set_repeated_text"),
    ("risk_catalyst", "risks"): ("slide.bullet_points", "set_repeated_text"),
    ("capability_map", "stages"): ("slide.bullet_points", "set_repeated_text"),
    ("chart_text", "takeaway"): ("slide.key_message", "set_text"),
    ("chart_text", "bullets"): ("slide.bullet_points", "set_bullets"),
    ("chart_text", "conclusion"): ("slide.key_message", "set_text"),
    ("valuation", "headline"): ("slide.key_message", "set_text"),
    ("valuation", "bullets"): ("slide.bullet_points", "set_bullets"),
    ("valuation", "conclusion"): ("slide.key_message", "set_text"),
    ("earnings_forecast", "note"): ("slide.key_message", "set_text"),
    ("earnings_forecast", "assumptions"): ("slide.bullet_points", "set_bullets"),
}

_IMAGE_ONLY_REMOVE_SHAPES = (
    "logic_chart_panel",
    "logic_chart_title",
    "logic_takeaway",
    "logic_main_point",
    "logic_bullets",
    "logic_conclusion",
    "logic_conclusion_text",
)


class TemplateProfileError(ValueError):
    """Raised when a Layout Map cannot produce a valid Template Profile."""


def _first_text_run(shape: Any) -> Any | None:
    if not getattr(shape, "has_text_frame", False):
        return None
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            if run.text.strip():
                return run
    return None


def _find_style_shape(presentation: Presentation, names: Sequence[str]) -> Any | None:
    wanted = {name.casefold() for name in names}
    for slide in presentation.slides:
        for shape in slide.shapes:
            if str(getattr(shape, "name", "")).casefold() in wanted:
                if _first_text_run(shape) is not None:
                    return shape
    return None


def _text_style_token(
    presentation: Presentation,
    names: Sequence[str],
    *,
    fallback_size: float,
    fallback_color: str,
    fallback_bold: bool,
) -> dict[str, Any]:
    shape = _find_style_shape(presentation, names)
    run = _first_text_run(shape) if shape is not None else None
    font = run.font if run is not None else None
    color = fallback_color
    if (
        font is not None
        and font.color.type == MSO_COLOR_TYPE.RGB
        and font.color.rgb is not None
    ):
        color = str(font.color.rgb)
    return {
        "font_family": (
            str(font.name)
            if font is not None and font.name
            else "Microsoft YaHei"
        ),
        "font_size_pt": (
            float(font.size.pt)
            if font is not None and font.size is not None
            else fallback_size
        ),
        "bold": (
            bool(font.bold)
            if font is not None and font.bold is not None
            else fallback_bold
        ),
        "color": color,
        "alignment": "left",
        "vertical_alignment": "middle" if "title" in names else "top",
    }


def _adaptive_style_tokens(presentation: Presentation) -> dict[str, Any]:
    """Induce reusable style tokens from named shapes in the supplied template."""

    title = _text_style_token(
        presentation,
        ("title",),
        fallback_size=27,
        fallback_color="102A43",
        fallback_bold=True,
    )
    key_message = _text_style_token(
        presentation,
        ("thesis", "logic_main_point", "company_positioning"),
        fallback_size=21,
        fallback_color=title["color"],
        fallback_bold=True,
    )
    body = _text_style_token(
        presentation,
        ("logic_1_body", "agenda_desc_1", "risk_1_body"),
        fallback_size=12.75,
        fallback_color="334155",
        fallback_bold=False,
    )
    # Adaptive pages have less decorative scaffolding than exact template pages.
    # Preserve the template typography but keep body copy at a readable floor.
    key_message["font_size_pt"] = max(
        14.0, min(16.0, float(key_message["font_size_pt"]))
    )
    body["font_size_pt"] = max(14.0, float(body["font_size_pt"]))
    table_font_size = max(8.0, min(10.0, float(body["font_size_pt"]) - 4.0))
    return {
        "slide_title": title,
        "key_message": key_message,
        "body_text": body,
        "primary_visual": {
            "font_family": body["font_family"],
            "font_size_pt": 10,
            "legend_font_size_pt": 9,
            "data_label_font_size_pt": 9,
            "text_color": body["color"],
            "axis_color": "94A3B8",
            "gridline_color": "E2E8F0",
            "series_colors": [
                "2F75B5",
                "ED7D31",
                "70AD47",
                "A5A5A5",
                "FFC000",
                "5B9BD5",
            ],
            "legend_position": "bottom",
            "line_width_pt": 2.25,
            "marker_size_pt": 6,
            "gap_width": 65,
        },
        "table": {
            "font_family": body["font_family"],
            "font_size_pt": table_font_size,
            "text_color": body["color"],
            "header_fill": "102A43",
            "header_text_color": "FFFFFF",
            "body_fill": "FFFFFF",
            "stripe_fill": "F5F7FA",
            "border_color": "D9E2EC",
            "border_width_pt": 0.6,
            "margin_horizontal_pt": 5,
            "margin_vertical_pt": 3,
            "header_height_pt": 30,
            "row_height_pt": 26,
            "first_column_bold": True,
            "max_rows": 18,
            "max_columns": 8,
        },
    }


def _load_json(path: Path, label: str) -> dict[str, Any]:
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
    except (FileNotFoundError, json.JSONDecodeError) as exc:
        raise TemplateProfileError(f"cannot load {label}: {path}: {exc}") from exc
    if not isinstance(value, dict):
        raise TemplateProfileError(f"{label} root must be an object")
    return value


def _target(value: object) -> dict[str, Any] | None:
    if not isinstance(value, Mapping):
        return None
    result = {
        key: value[key]
        for key in ("name", "shape_id")
        if key in value
    }
    return result or None


def _bounds_target(value: object) -> dict[str, Any] | None:
    if not isinstance(value, Mapping):
        return None
    if not all(isinstance(value.get(key), (int, float)) for key in ("left", "top", "width", "height")):
        return None
    return {
        "bounds_in": {
            key: value[key] for key in ("left", "top", "width", "height")
        }
    }


def _binding_source(layout_id: str, field_name: str) -> tuple[str, str] | None:
    return _SOURCE_BINDINGS.get((layout_id, field_name)) or _SOURCE_BINDINGS.get(
        ("*", field_name)
    )


def _binding(
    layout_id: str, field_name: str, spec: Mapping[str, Any]
) -> dict[str, Any] | None:
    source_and_operation = _binding_source(layout_id, field_name)
    if source_and_operation is None:
        return None
    source, operation = source_and_operation
    result: dict[str, Any] = {
        "binding_id": field_name,
        "kind": {
            "set_text": "auto" if source.startswith("compiler.") else "text",
            "set_bullets": "bullet_list",
            "set_repeated_text": "repeated_text",
        }[operation],
        "operation": operation,
        "source": source,
        "required": bool(spec.get("required", True)),
    }
    if operation == "set_repeated_text":
        targets = []
        for group in spec.get("targets", []):
            if not isinstance(group, Mapping):
                continue
            resolved = {
                str(name): target
                for name, raw_target in group.items()
                if (target := _target(raw_target)) is not None
            }
            if resolved:
                targets.append(resolved)
        if not targets:
            return None
        result["targets"] = targets
        result["max_items"] = int(spec.get("max_items", len(targets)))
    else:
        target = _target(spec.get("target"))
        if target is None:
            return None
        result["target"] = target
        if field_name == "source_refs":
            result["separator"] = ", "
    return result


def _slot(field_name: str, spec: Mapping[str, Any]) -> dict[str, Any] | None:
    field_type = spec.get("type")
    if field_type == "chart_slot":
        kind = "chart"
        operation = "render_chart"
        target = _target(spec.get("anchor"))
    elif field_type == "table":
        kind = "table"
        operation = "render_table"
        target = _target(spec.get("target"))
    elif field_type == "image_slot":
        kind = "image"
        operation = "render_image"
        target = _bounds_target(spec.get("bounds_in"))
    else:
        return None
    if target is None:
        return None
    capacity: dict[str, int] = {"max_items": 1}
    for source_key, target_key in (
        ("max_rows", "max_rows"),
        ("max_columns", "max_columns"),
        ("max_categories", "max_categories"),
        ("max_series", "max_series"),
    ):
        value = spec.get(source_key)
        if isinstance(value, int) and value > 0:
            capacity[target_key] = value
    result: dict[str, Any] = {
        "slot_id": field_name,
        "kind": kind,
        "operation": operation,
        "required": bool(spec.get("required", True)),
        "target": target,
        "capacity": capacity,
    }
    title_target = _target(spec.get("title_target"))
    if title_target is not None:
        result["title_target"] = title_target
    return result


def _resolution(layout_map: Mapping[str, Any]) -> dict[str, Any]:
    source = layout_map.get("layout_resolution", {})
    if not isinstance(source, Mapping):
        raise TemplateProfileError("Layout Map layout_resolution must be an object")
    return {
        "page_role_overrides": dict(source.get("page_role_overrides", {})),
        "semantic_visual_rules": list(source.get("semantic_visual_rules", [])),
        "slide_type_defaults": dict(source.get("slide_type_defaults", {})),
    }


def build_template_profile(
    layout_map: Mapping[str, Any],
    template_path: Path,
    *,
    profile_id: str = "financial-report-v1",
) -> dict[str, Any]:
    """Convert the legacy Layout Map into the explicit new contract."""

    if not template_path.is_file():
        raise TemplateProfileError(f"template not found: {template_path}")
    presentation = layout_map.get("presentation", {})
    layouts = layout_map.get("layouts")
    if not isinstance(presentation, Mapping) or not isinstance(layouts, Mapping):
        raise TemplateProfileError("Layout Map must contain presentation and layouts")
    template_presentation = Presentation(str(template_path))
    profile_layouts: dict[str, Any] = {}
    for layout_id, raw_layout in layouts.items():
        if not isinstance(raw_layout, Mapping):
            continue
        fields = raw_layout.get("fields", {})
        if not isinstance(fields, Mapping):
            fields = {}
        bindings = [
            value
            for name, spec in fields.items()
            if isinstance(spec, Mapping)
            and (value := _binding(str(layout_id), str(name), spec)) is not None
        ]
        slots = [
            value
            for name, spec in fields.items()
            if isinstance(spec, Mapping)
            and (value := _slot(str(name), spec)) is not None
        ]
        layout = {
            "template_slide": raw_layout.get("template_slide"),
            "description": str(raw_layout.get("description", "")),
            "bindings": bindings,
            "slots": slots,
        }
        bound_names = {
            str(target.get("name"))
            for binding in [*bindings, *slots]
            for target in (
                [binding.get("target")]
                if isinstance(binding.get("target"), Mapping)
                else [
                    nested
                    for group in binding.get("targets", [])
                    if isinstance(group, Mapping)
                    for nested in group.values()
                    if isinstance(nested, Mapping)
                ]
            )
            if isinstance(target, Mapping) and target.get("name")
        }
        template_slide = int(raw_layout.get("template_slide", 0))
        if 1 <= template_slide <= len(template_presentation.slides):
            layout["has_unbound_example_content"] = any(
                shape.name.startswith("metric_") and shape.name not in bound_names
                for shape in template_presentation.slides[template_slide - 1].shapes
            )
        if layout_id == "image_only_layout":
            layout["remove_shapes"] = list(_IMAGE_ONLY_REMOVE_SHAPES)
        profile_layouts[str(layout_id)] = layout
    blank_layout_index = min(
        range(len(template_presentation.slide_layouts)),
        key=lambda index: (
            len(template_presentation.slide_layouts[index].placeholders),
            index,
        ),
    )
    width_in = float(presentation.get("width_in", 13.333))
    height_in = float(presentation.get("height_in", 7.5))
    style_tokens = _adaptive_style_tokens(template_presentation)
    return {
        "schema_version": "1.0.0",
        "profile_id": profile_id,
        "template": {
            "file": template_path.name,
            "sha256": hashlib.sha256(template_path.read_bytes()).hexdigest(),
            "slide_count": presentation.get("slide_count"),
            "width_in": presentation.get("width_in"),
            "height_in": presentation.get("height_in"),
        },
        "supported_operations": list(SUPPORTED_OPERATIONS),
        "layout_resolution": _resolution(layout_map),
        "adaptive_canvas": {
            "enabled": True,
            "base": {
                "mode": "blank",
                "slide_layout_index": blank_layout_index,
                "clear_placeholders": True,
                "background_color": "FFFFFF",
            },
            "safe_area_in": {
                "left": round(width_in * 0.04, 4),
                "top": round(height_in * 0.04, 4),
                "width": round(width_in * 0.92, 4),
                "height": round(height_in * 0.92, 4),
            },
            "style_tokens": {
                **style_tokens,
            },
        },
        "layouts": profile_layouts,
    }


def validate_template_profile(
    profile: Mapping[str, Any], schema: Mapping[str, Any]
) -> list[str]:
    errors = sorted(
        Draft202012Validator(schema).iter_errors(profile),
        key=lambda error: list(error.absolute_path),
    )
    messages = [error.message for error in errors]
    layouts = profile.get("layouts", {})
    if isinstance(layouts, Mapping):
        known = set(layouts)
        resolution = profile.get("layout_resolution", {})
        if isinstance(resolution, Mapping):
            targets = [
                *dict(resolution.get("page_role_overrides", {})).values(),
                *dict(resolution.get("slide_type_defaults", {})).values(),
                *[
                    rule.get("layout_id")
                    for rule in resolution.get("semantic_visual_rules", [])
                    if isinstance(rule, Mapping)
                ],
            ]
            for target in targets:
                if target not in known:
                    messages.append(f"unknown layout resolution target: {target!r}")
        for layout_id, layout in layouts.items():
            if not isinstance(layout, Mapping):
                continue
            slot_ids = [
                slot.get("slot_id")
                for slot in layout.get("slots", [])
                if isinstance(slot, Mapping)
            ]
            if len(slot_ids) != len(set(slot_ids)):
                messages.append(f"layout {layout_id!r} contains duplicate slot_id")
    return messages


def parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Build or validate a deterministic Template Profile"
    )
    parser.add_argument("layout_map", type=Path)
    parser.add_argument("template", type=Path)
    parser.add_argument("-o", "--output", type=Path, required=True)
    parser.add_argument("--schema", type=Path, default=DEFAULT_SCHEMA)
    parser.add_argument("--profile-id", default="financial-report-v1")
    return parser.parse_args(argv)


def main(argv: Sequence[str] | None = None) -> int:
    args = parse_args(argv)
    try:
        layout_map = _load_json(args.layout_map, "Layout Map")
        schema = _load_json(args.schema, "Template Profile schema")
        profile = build_template_profile(
            layout_map, args.template, profile_id=args.profile_id
        )
        errors = validate_template_profile(profile, schema)
        if errors:
            raise TemplateProfileError("; ".join(errors))
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(
            json.dumps(profile, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )
    except (OSError, TemplateProfileError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1
    print(f"Created Template Profile: {args.output}")
    print(f"Layouts: {len(profile['layouts'])}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

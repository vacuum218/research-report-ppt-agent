from __future__ import annotations

import json
from pathlib import Path
from types import SimpleNamespace

import pytest
from jsonschema import Draft202012Validator
from pptx import Presentation
from pptx.enum.chart import (
    XL_CHART_TYPE,
    XL_LEGEND_POSITION,
    XL_MARKER_STYLE,
)
from pptx.oxml.ns import qn
from pptx.util import Inches

from ppt_engine.visualization_renderer import (
    VisualizationRenderError,
    render_chart,
)


PROJECT_ROOT = Path(__file__).resolve().parents[2]
CHART_EXAMPLES = (
    "week3_t3_4_line.json",
    "week3_t3_4_column.json",
    "week3_t3_4_bar.json",
    "week3_t3_4_pie.json",
)


def _anchor() -> SimpleNamespace:
    return SimpleNamespace(
        left=Inches(0.8),
        top=Inches(1.2),
        width=Inches(8.4),
        height=Inches(5.2),
    )


def _slide():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    return presentation, slide


def _chart(
    chart_type: str,
    *,
    values: list[float | None] | None = None,
) -> dict:
    categories = ["2022A", "2023A", "2024A", "2025E"]
    first_values = values or [80, 100, 125, 150]
    series = [
        {
            "name": "主营业务",
            "values": first_values,
            "number_format": "0.0",
        }
    ]
    if chart_type != "pie":
        series.append(
            {
                "name": "第二业务",
                "values": [40, 55, 70, 90],
                "number_format": "0.0",
            }
        )
    return {
        "chart_type": chart_type,
        "title": f"{chart_type} 验收",
        "unit": "亿元",
        "categories": categories,
        "series": series,
        "forecast_start_index": 3 if chart_type == "line" else None,
        "source_refs": ["src_week3_t3_4"],
    }


@pytest.mark.parametrize(
    ("chart_type", "expected_type"),
    [
        ("line", XL_CHART_TYPE.LINE_MARKERS),
        ("column", XL_CHART_TYPE.COLUMN_CLUSTERED),
        ("bar", XL_CHART_TYPE.BAR_CLUSTERED),
        ("pie", XL_CHART_TYPE.PIE),
    ],
)
def test_native_chart_types_are_reopenable(
    tmp_path,
    chart_type,
    expected_type,
):
    presentation, slide = _slide()
    render_chart(slide, _anchor(), _chart(chart_type))
    output = tmp_path / f"{chart_type}.pptx"
    presentation.save(output)

    reopened = Presentation(output)
    rendered = next(
        shape.chart
        for shape in reopened.slides[0].shapes
        if shape.has_chart
    )

    assert rendered.chart_type == expected_type
    assert rendered.has_title is False
    assert rendered.has_legend is True
    assert rendered.legend.position == XL_LEGEND_POSITION.BOTTOM
    assert len(rendered.series) == (1 if chart_type == "pie" else 2)


def test_line_chart_preserves_null_and_marks_forecast_points(tmp_path):
    presentation, slide = _slide()
    visualization = _chart("line", values=[80, None, 125, 150])
    visualization["forecast_start_index"] = 2

    render_chart(slide, _anchor(), visualization)
    output = tmp_path / "line-with-null-and-forecast.pptx"
    presentation.save(output)
    reopened = Presentation(output)
    chart = next(
        shape.chart
        for shape in reopened.slides[0].shapes
        if shape.has_chart
    )

    cached_points = chart._chartSpace.xpath(
        ".//c:lineChart/c:ser[1]/c:val/c:numRef/c:numCache/c:pt"
    )
    assert {
        point.get("idx"): point.xpath("./c:v")[0].text
        for point in cached_points
    } == {"0": "80", "2": "125", "3": "150"}
    assert (
        chart._chartSpace.chart.find(qn("c:dispBlanksAs")).get("val")
        == "span"
    )
    assert chart.series[0].marker.style == XL_MARKER_STYLE.CIRCLE
    forecast_points = chart._chartSpace.xpath(
        ".//c:lineChart/c:ser[1]/c:dPt/c:idx"
    )
    assert {point.get("val") for point in forecast_points} == {"2", "3"}


def test_type_specific_chart_style_is_applied():
    _, column_slide = _slide()
    column = render_chart(column_slide, _anchor(), _chart("column"))
    assert column.plots[0].gap_width == 65
    assert str(column.series[0].format.fill.fore_color.rgb) == "2F75B5"
    assert column.value_axis.has_major_gridlines is True

    _, bar_slide = _slide()
    bar = render_chart(bar_slide, _anchor(), _chart("bar"))
    assert bar.plots[0].has_data_labels is True
    assert bar.plots[0].data_labels.show_value is True

    _, pie_slide = _slide()
    pie = render_chart(pie_slide, _anchor(), _chart("pie"))
    assert pie.plots[0].has_data_labels is True
    assert pie.plots[0].data_labels.show_percentage is True
    assert str(pie.series[0].points[0].format.fill.fore_color.rgb) == "2F75B5"


@pytest.mark.parametrize(
    ("mutation", "message"),
    [
        (
            lambda value: value["series"].append(
                {"name": "非法第二系列", "values": [1, 2, 3, 4]}
            ),
            "exactly one series",
        ),
        (
            lambda value: value["series"][0]["values"].__setitem__(1, None),
            "missing values",
        ),
        (
            lambda value: value["series"][0]["values"].__setitem__(1, -1),
            "negative values",
        ),
    ],
)
def test_pie_rejects_misleading_data(mutation, message):
    _, slide = _slide()
    visualization = _chart("pie")
    mutation(visualization)

    with pytest.raises(VisualizationRenderError, match=message):
        render_chart(slide, _anchor(), visualization)


def test_fixed_chart_examples_are_schema_valid():
    schema = json.loads(
        (
            PROJECT_ROOT / "schemas/visualization.schema.json"
        ).read_text(encoding="utf-8")
    )
    validator = Draft202012Validator(schema)

    for filename in CHART_EXAMPLES:
        value = json.loads(
            (
                PROJECT_ROOT
                / "examples/visualization_charts"
                / filename
            ).read_text(encoding="utf-8")
        )
        validator.validate(value)

from __future__ import annotations

import json
from pathlib import Path
from types import SimpleNamespace

import pytest
from jsonschema import Draft202012Validator
from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches

from ppt_engine.visualization_renderer import (
    VisualizationRenderError,
    render_table,
)


PROJECT_ROOT = Path(__file__).resolve().parents[2]
TABLE_EXAMPLES = (
    "week3_t3_5_forecast.json",
    "week3_t3_5_valuation.json",
    "week3_t3_5_long_labels.json",
    "week3_t3_5_capacity_boundary.json",
)


def _anchor() -> SimpleNamespace:
    return SimpleNamespace(
        left=Inches(0.75),
        top=Inches(1.25),
        width=Inches(11.83),
        height=Inches(5.45),
    )


def _slide():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    return presentation, slide


def _table() -> dict:
    return {
        "title": "盈利预测",
        "unit": "亿元",
        "columns": ["指标", "2024A", "2025E", "备注"],
        "rows": [
            ["营业收入", 100, 120, "稳健增长"],
            ["归母净利润", 8.5, 10.2, "利润率改善"],
            ["毛利率", "28.0%", None, "待披露"],
        ],
        "source_refs": ["src_week3_t3_5"],
    }


def test_native_table_style_and_data_survive_reopen(tmp_path):
    presentation, slide = _slide()
    rendered = render_table(slide, _anchor(), _table())
    output = tmp_path / "styled-table.pptx"
    presentation.save(output)

    reopened = Presentation(output)
    shape = next(
        item for item in reopened.slides[0].shapes if item.has_table
    )
    table = shape.table

    assert shape.name == "compiled_table"
    assert len(table.rows) == 4
    assert len(table.columns) == 4
    assert table.cell(0, 0).text == "指标"
    assert table.cell(3, 2).text == ""
    assert str(table.cell(0, 0).fill.fore_color.rgb) == "102A43"
    assert str(table.cell(0, 0).text_frame.paragraphs[0].runs[0].font.color.rgb) == (
        "FFFFFF"
    )
    assert table.cell(0, 0).text_frame.paragraphs[0].runs[0].font.bold is True
    assert str(table.cell(2, 0).fill.fore_color.rgb) == "F5F7FA"
    assert table.cell(1, 1).text_frame.paragraphs[0].alignment == PP_ALIGN.CENTER
    assert table.cell(3, 1).text_frame.paragraphs[0].alignment == PP_ALIGN.CENTER
    assert table.cell(1, 0).text_frame.paragraphs[0].alignment == PP_ALIGN.CENTER
    assert table.cell(1, 0).text_frame.paragraphs[0].runs[0].font.bold is True
    assert table.cell(1, 0).vertical_anchor == MSO_ANCHOR.MIDDLE
    assert sum(column.width for column in table.columns) == shape.width
    assert table.columns[0].width > table.columns[1].width
    assert table.rows[0].height > table.rows[1].height

    left_border = table.cell(1, 1)._tc.get_or_add_tcPr().find(qn("a:lnL"))
    assert left_border is not None
    border_rgb = left_border.find(qn("a:solidFill")).find(qn("a:srgbClr"))
    assert border_rgb.get("val") == "D9E2EC"


@pytest.mark.parametrize(
    ("value", "message"),
    [
        (
            {
                "columns": ["项目", "数值"],
                "rows": [[f"项目{index}", index] for index in range(19)],
            },
            "max_rows=18",
        ),
        (
            {
                "columns": [f"列{index}" for index in range(9)],
                "rows": [[index for index in range(9)]],
            },
            "max_columns=8",
        ),
        (
            {
                "columns": ["项目", "数值"],
                "rows": [["缺少数值"]],
            },
            "must match columns length",
        ),
    ],
)
def test_table_renderer_rejects_overflow_or_inconsistent_rows(value, message):
    _, slide = _slide()

    with pytest.raises(VisualizationRenderError, match=message):
        render_table(slide, _anchor(), value)


def test_fixed_table_examples_are_schema_valid():
    schema = json.loads(
        (
            PROJECT_ROOT / "schemas/visualization.schema.json"
        ).read_text(encoding="utf-8")
    )
    validator = Draft202012Validator(schema)

    for filename in TABLE_EXAMPLES:
        value = json.loads(
            (
                PROJECT_ROOT
                / "examples/visualization_tables"
                / filename
            ).read_text(encoding="utf-8")
        )
        validator.validate(value)

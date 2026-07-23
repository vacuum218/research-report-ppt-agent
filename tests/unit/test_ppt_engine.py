from __future__ import annotations

import base64
import json
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn
from pptx.util import Inches

from ppt_engine.renderer import RenderError, render_presentation


PROJECT_ROOT = Path(__file__).resolve().parents[2]


def load(relative: str):
    return json.loads((PROJECT_ROOT / relative).read_text(encoding="utf-8"))


def image_visualization(asset_path: str) -> dict:
    return {
        "type": "image",
        "title": "原始研报图片",
        "source": {"kind": "figure", "id": "fig-001"},
        "asset_path": asset_path,
        "source_refs": ["src_annual_2025"],
        "sources": [{"kind": "figure", "id": "fig-001"}],
    }


def write_test_png(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(
        base64.b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAIAAAABCAQAAABeK7cBAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
        )
    )


def test_outline_only_renderer_creates_reopenable_pptx(tmp_path):
    outline = load("examples/slide_outline_valid.json")
    layout_map = load("templates/template_layout_map.json")
    output = tmp_path / "outline-only.pptx"

    render_presentation(
        outline,
        layout_map,
        PROJECT_ROOT / "templates/financial_report_template_v1.pptx",
        output,
    )

    presentation = Presentation(output)
    assert len(presentation.slides) == len(outline["slides"])
    assert presentation.slides[0].shapes[2].name == "cover_title"
    assert presentation.slides[0].element.cSld.find(qn("p:bg")) is not None
    visual_slide = presentation.slides[1]
    assert not any(shape.has_chart for shape in visual_slide.shapes)
    assert not any(shape.has_table for shape in visual_slide.shapes)


def test_outline_only_renderer_replaces_template_table_with_empty_anchor(tmp_path):
    outline = load("examples/generated/002544_2025-10-28_slide_outline.json")
    layout_map = load("templates/template_layout_map.json")
    output = tmp_path / "no-template-sample-data.pptx"

    render_presentation(
        outline,
        layout_map,
        PROJECT_ROOT / "templates/financial_report_template_v1.pptx",
        output,
    )

    forecast_slide = Presentation(output).slides[6]
    assert not any(shape.has_table for shape in forecast_slide.shapes)
    assert any(shape.name == "forecast_table" for shape in forecast_slide.shapes)


def test_renderer_supports_chart_and_table_bindings(tmp_path):
    outline = load("examples/slide_outline_valid.json")
    layout_map = load("templates/template_layout_map.json")
    chart = load("examples/visualization_valid.json")
    table = load("examples/visualization_table_valid.json")
    output = tmp_path / "visuals.pptx"

    render_presentation(
        outline,
        layout_map,
        PROJECT_ROOT / "templates/financial_report_template_v1.pptx",
        output,
        visualizations_by_slide={"slide_002": [chart, table]},
    )

    presentation = Presentation(output)
    visual_slide = presentation.slides[1]
    assert any(shape.has_chart for shape in visual_slide.shapes)
    forecast_table = next(shape for shape in visual_slide.shapes if shape.name == "forecast_table")
    assert forecast_table.table.cell(0, 0).text == "项目"
    assert forecast_table.table.cell(1, 0).text == "营业收入"


def test_renderer_rejects_unknown_visualization_slide(tmp_path):
    outline = load("examples/slide_outline_valid.json")
    layout_map = load("templates/template_layout_map.json")
    chart = load("examples/visualization_valid.json")

    with pytest.raises(RenderError, match="unknown slide_id"):
        render_presentation(
            outline,
            layout_map,
            PROJECT_ROOT / "templates/financial_report_template_v1.pptx",
            tmp_path / "invalid.pptx",
            visualizations_by_slide={"slide_missing": [chart]},
        )


def test_renderer_inserts_image_shape_from_safe_relative_path(tmp_path, monkeypatch):
    outline = load("examples/slide_outline_valid.json")
    layout_map = load("templates/template_layout_map.json")
    image_path = tmp_path / "assets" / "figures" / "fig-001.png"
    write_test_png(image_path)
    monkeypatch.chdir(tmp_path)
    output = tmp_path / "image.pptx"

    render_presentation(
        outline,
        layout_map,
        PROJECT_ROOT / "templates/financial_report_template_v1.pptx",
        output,
        visualizations_by_slide={
            "slide_002": [image_visualization("assets/figures/fig-001.png")]
        },
    )

    presentation = Presentation(output)
    pictures = [
        shape
        for shape in presentation.slides[1].shapes
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
    ]
    assert len(pictures) == 1
    assert pictures[0].name == "image_fig-001"
    assert pictures[0].width > Inches(3)


def test_figure_page_uses_full_width_image_only_layout_and_asset_root(tmp_path):
    outline = load("examples/slide_outline_valid.json")
    layout_map = load("templates/template_layout_map.json")
    slide = outline["slides"][1]
    slide["slide_type"] = "figure_page"
    slide["title"] = "公司 PDT 星通站实现一站多能"
    slide["bullet_points"] = []
    slide["evidence_refs"] = [{"kind": "figure", "id": "fig-007"}]
    slide["visual_candidates"] = []
    bundle = tmp_path / "document_bundle"
    image_path = bundle / "assets" / "figures" / "fig-007.png"
    write_test_png(image_path)
    output = tmp_path / "figure-page.pptx"
    figure_image = image_visualization("assets/figures/fig-007.png")
    figure_image["source"]["id"] = "fig-007"
    figure_image["sources"][0]["id"] = "fig-007"

    render_presentation(
        outline,
        layout_map,
        PROJECT_ROOT / "templates/financial_report_template_v1.pptx",
        output,
        visualizations_by_slide={
            "slide_002": [figure_image]
        },
        asset_root=bundle,
    )

    presentation = Presentation(output)
    pictures = [
        shape
        for shape in presentation.slides[1].shapes
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
    ]
    assert len(pictures) == 1
    assert pictures[0].name == "image_fig-007"
    assert pictures[0].width > Inches(10)
    assert not {
        "logic_chart_panel",
        "logic_takeaway",
        "logic_conclusion",
    } & {shape.name for shape in presentation.slides[1].shapes}


def test_renderer_fails_when_image_asset_does_not_exist(tmp_path, monkeypatch):
    outline = load("examples/slide_outline_valid.json")
    layout_map = load("templates/template_layout_map.json")
    monkeypatch.chdir(tmp_path)

    with pytest.raises(RenderError, match="image asset does not exist"):
        render_presentation(
            outline,
            layout_map,
            PROJECT_ROOT / "templates/financial_report_template_v1.pptx",
            tmp_path / "missing.pptx",
            visualizations_by_slide={
                "slide_002": [image_visualization("assets/figures/missing.png")]
            },
        )


def test_renderer_rejects_image_path_escape(tmp_path):
    outline = load("examples/slide_outline_valid.json")
    layout_map = load("templates/template_layout_map.json")

    with pytest.raises(RenderError, match="schema validation failed"):
        render_presentation(
            outline,
            layout_map,
            PROJECT_ROOT / "templates/financial_report_template_v1.pptx",
            tmp_path / "escape.pptx",
            visualizations_by_slide={
                "slide_002": [image_visualization("../../outside.png")]
            },
        )

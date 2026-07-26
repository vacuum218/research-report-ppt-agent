from __future__ import annotations

import json
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.enum.text import MSO_AUTO_SIZE

from ppt_engine.slide_builder import SlideBuildError, execute_compiled_operations


PROJECT_ROOT = Path(__file__).resolve().parents[2]


def test_executor_sets_only_explicit_text_operations(tmp_path):
    prs = Presentation(PROJECT_ROOT / "templates/financial_report_template_v1.pptx")
    slide = prs.slides[2]
    title = next(shape for shape in slide.shapes if shape.name == "title")

    execute_compiled_operations(
        slide,
        [
            {
                "op": "set_text",
                "binding_id": "title",
                "target": {"name": "title"},
                "value": "确定性标题",
            }
        ],
        visualizations_by_id={},
        asset_root=tmp_path,
    )

    assert title.text == "确定性标题"
    assert title.text_frame.auto_size == MSO_AUTO_SIZE.NONE


def test_executor_rejects_missing_explicit_target(tmp_path):
    prs = Presentation(PROJECT_ROOT / "templates/financial_report_template_v1.pptx")

    with pytest.raises(SlideBuildError, match="compiled target is missing"):
        execute_compiled_operations(
            prs.slides[0],
            [
                {
                    "op": "set_text",
                    "binding_id": "title",
                    "target": {"name": "missing"},
                    "value": "内容",
                }
            ],
            visualizations_by_id={},
            asset_root=tmp_path,
        )


def test_executor_rejects_unknown_operation(tmp_path):
    prs = Presentation(PROJECT_ROOT / "templates/financial_report_template_v1.pptx")

    with pytest.raises(SlideBuildError, match="unsupported compiled operation"):
        execute_compiled_operations(
            prs.slides[0],
            [{"op": "guess_layout"}],
            visualizations_by_id={},
            asset_root=tmp_path,
        )


def test_adaptive_text_box_disables_powerpoint_auto_shrink(tmp_path):
    prs = Presentation(PROJECT_ROOT / "templates/financial_report_template_v1.pptx")
    slide = prs.slides[0]

    execute_compiled_operations(
        slide,
        [
            {
                "op": "add_text_box",
                "element_id": "body",
                "target": {
                    "bounds_in": {
                        "left": 1,
                        "top": 1,
                        "width": 5,
                        "height": 2,
                    }
                },
                "value": "分页后的正文",
                "style": {
                    "font_family": "Microsoft YaHei",
                    "font_size_pt": 14,
                    "bold": False,
                    "color": "000000",
                    "alignment": "left",
                    "vertical_alignment": "top",
                },
            }
        ],
        visualizations_by_id={},
        asset_root=tmp_path,
    )

    shape = slide.shapes[-1]
    assert shape.text_frame.auto_size == MSO_AUTO_SIZE.NONE
    assert shape.text_frame.paragraphs[0].runs[0].font.size.pt == 14

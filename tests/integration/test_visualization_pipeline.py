from __future__ import annotations

import json
from pathlib import Path

from pptx import Presentation

from document_bundle.markdown import build_from_markdown
from document_intelligence import load_document_intelligence
from ppt_engine.renderer import render_presentation
from visualization_generator.generate_visualizations import (
    bindings_from_artifacts,
    generate_visualizations,
    preflight_visualizations,
)


PROJECT_ROOT = Path(__file__).resolve().parents[2]


def _load(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def test_real_report_to_visualization_to_pptx_pipeline(tmp_path):
    bundle = tmp_path / "document_bundle"
    build_from_markdown(PROJECT_ROOT / "data/reports/agent/002544_2025-10-28.md", bundle)
    document = load_document_intelligence(bundle, PROJECT_ROOT / "schemas/document_bundle.schema.json")
    outline = _load(PROJECT_ROOT / "examples/generated/002544_2025-10-28_slide_outline.json")
    layout_map = _load(PROJECT_ROOT / "templates/template_layout_map.json")

    artifacts, issues = generate_visualizations(outline, document)
    bindings = bindings_from_artifacts(artifacts)

    assert not issues
    assert "chart_type" in bindings["slide_006"][0]
    assert "columns" in bindings["slide_007"][0]
    assert "columns" in bindings["slide_009"][0]
    assert preflight_visualizations(outline, layout_map, bindings) == []

    output = tmp_path / "full-pipeline.pptx"
    render_presentation(
        outline,
        layout_map,
        PROJECT_ROOT / "templates/financial_report_template_v1.pptx",
        output,
        visualizations_by_slide=bindings,
    )

    presentation = Presentation(output)
    assert len(presentation.slides) == len(outline["slides"])
    assert any(shape.has_chart for shape in presentation.slides[5].shapes)
    assert any(shape.has_table for shape in presentation.slides[6].shapes)
    assert any(shape.has_table for shape in presentation.slides[8].shapes)

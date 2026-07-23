from __future__ import annotations

import json
from pathlib import Path

from pptx import Presentation

from ppt_engine.compiler import compile_layout_plan
from ppt_engine.renderer import render_compiled_plan
from tools.build_template_profile import build_template_profile
from visualization_generator.manifest import (
    canonical_sha256,
    load_visualization_manifest,
)


PROJECT_ROOT = Path(__file__).resolve().parents[2]


def load(relative: str) -> dict:
    return json.loads((PROJECT_ROOT / relative).read_text(encoding="utf-8"))


def write_json(path: Path, value: dict) -> None:
    path.write_text(json.dumps(value, ensure_ascii=False), encoding="utf-8")


def test_compiled_plan_renders_existing_chart_and_table_effect(tmp_path):
    outline = load("examples/slide_outline_valid.json")
    chart = load("examples/visualization_valid.json")
    table = load("examples/visualization_table_valid.json")
    write_json(tmp_path / "chart.json", chart)
    write_json(tmp_path / "table.json", table)
    bundle = tmp_path / "bundle"
    bundle.mkdir()
    write_json(
        tmp_path / "manifest.json",
        {
            "schema_version": "3.0.0",
            "outline_sha256": canonical_sha256(outline),
            "document_source_sha256": "a" * 64,
            "asset_root": str(bundle),
            "bindings": [
                {
                    "slide_id": "slide_002",
                    "visualization_id": "visual_chart",
                    "visual_type": "chart",
                    "sources": [{"kind": "table", "id": "table-001"}],
                    "visualization_file": "chart.json",
                },
                {
                    "slide_id": "slide_002",
                    "visualization_id": "visual_table",
                    "visual_type": "table",
                    "sources": [{"kind": "table", "id": "table-001"}],
                    "visualization_file": "table.json",
                },
            ],
        },
    )
    profile = build_template_profile(
        load("templates/template_layout_map.json"),
        PROJECT_ROOT / "templates/financial_report_template_v1.pptx",
    )
    plan = compile_layout_plan(
        outline,
        profile,
        load_visualization_manifest(tmp_path / "manifest.json"),
    )
    output = tmp_path / "compiled.pptx"

    render_compiled_plan(
        plan,
        template_path=PROJECT_ROOT / "templates/financial_report_template_v1.pptx",
        output_path=output,
    )

    prs = Presentation(output)
    assert len(prs.slides) == 2
    assert any(shape.has_chart for shape in prs.slides[1].shapes)
    table_shape = next(shape for shape in prs.slides[1].shapes if shape.has_table)
    assert table_shape.table.cell(0, 0).text == "项目"
    assert next(shape for shape in prs.slides[1].shapes if shape.name == "title").text == "未来盈利预测"

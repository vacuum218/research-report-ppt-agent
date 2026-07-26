from __future__ import annotations

import json
from copy import deepcopy
from dataclasses import replace
from pathlib import Path

import pytest
from pptx import Presentation

from document_intelligence import load_document_intelligence
from tools.build_week3_t3_6_end_to_end import build_acceptance_output
from visualization_generator.audit import (
    NumericAuditError,
    audit_visualization_artifacts,
)
from visualization_generator.generate_visualizations import generate_visualizations
from visualization_generator.numeric_facts import build_numeric_fact_ledger
from visualization_generator.planning import plan_visualizations


PROJECT_ROOT = Path(__file__).resolve().parents[2]
FIXTURE_ROOT = PROJECT_ROOT / "examples/week3_t3_6"
BUNDLE_PATH = FIXTURE_ROOT / "document_bundle"
OUTLINE_PATH = FIXTURE_ROOT / "outline.json"
BUNDLE_SCHEMA = PROJECT_ROOT / "schemas/document_bundle.schema.json"


def _load(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def test_fixed_scenarios_are_discovered_rejected_and_audited():
    outline = _load(OUTLINE_PATH)
    snapshot = load_document_intelligence(BUNDLE_PATH, BUNDLE_SCHEMA)

    plans = plan_visualizations(outline, snapshot)
    artifacts, issues = generate_visualizations(outline, snapshot)
    artifacts_by_slide = {artifact.slide_id: artifact for artifact in artifacts}

    assert not issues
    assert {plan.slide_id for plan in plans} == {
        "slide_t3_6_trend",
        "slide_t3_6_table",
        "slide_t3_6_composition",
    }
    assert set(artifacts_by_slide) == {
        "slide_t3_6_trend",
        "slide_t3_6_table",
        "slide_t3_6_composition",
    }
    assert artifacts_by_slide["slide_t3_6_trend"].data["chart_type"] == "bar"
    assert "columns" in artifacts_by_slide["slide_t3_6_table"].data
    assert (
        artifacts_by_slide["slide_t3_6_composition"].data["chart_type"]
        == "pie"
    )

    ledger = build_numeric_fact_ledger(snapshot)
    audit = audit_visualization_artifacts(artifacts, ledger)
    assert audit["status"] == "passed"
    assert audit["summary"] == {
        "visualization_count": 3,
        "audited_value_count": 10,
        "mismatch_count": 0,
    }

    trend = artifacts_by_slide["slide_t3_6_trend"]
    tampered_data = deepcopy(trend.data)
    tampered_data["series"][0]["values"][0] = 999
    tampered = replace(trend, data=tampered_data)
    with pytest.raises(NumericAuditError, match="numeric audit mismatch"):
        audit_visualization_artifacts(
            [
                tampered,
                artifacts_by_slide["slide_t3_6_table"],
                artifacts_by_slide["slide_t3_6_composition"],
            ],
            ledger,
        )


def test_acceptance_builder_writes_compiled_native_pptx(tmp_path):
    outputs = build_acceptance_output(output_directory=tmp_path)

    expected_keys = {
        "ledger",
        "audit",
        "scenarios",
        "manifest",
        "profile",
        "compiled_plan",
        "presentation",
    }
    assert set(outputs) == expected_keys
    assert all(path.is_file() for path in outputs.values())

    audit = _load(outputs["audit"])
    scenarios = _load(outputs["scenarios"])
    compiled_plan = _load(outputs["compiled_plan"])
    assert audit["status"] == "passed"
    assert scenarios["status"] == "passed"
    assert scenarios["successful_count"] == 3
    assert scenarios["rejected_count"] == 2
    assert len(compiled_plan["visualizations"]) == 3

    presentation = Presentation(outputs["presentation"])
    native_chart_count = sum(
        shape.has_chart
        for slide in presentation.slides
        for shape in slide.shapes
    )
    native_table_count = sum(
        shape.has_table
        for slide in presentation.slides
        for shape in slide.shapes
    )
    assert len(presentation.slides) == 5
    assert native_chart_count == 2
    assert native_table_count == 1

from __future__ import annotations

import hashlib
import json
import subprocess
import sys
from copy import deepcopy
from pathlib import Path

from pptx import Presentation

from tools.build_template_profile import build_template_profile
from pipeline_runner.run_pipeline import _partition_generation_issues
from visualization_generator.generator import GenerationIssue


PROJECT_ROOT = Path(__file__).resolve().parents[2]
FIXTURE_ROOT = PROJECT_ROOT / "examples/week3_t3_6"
BUNDLE_PATH = FIXTURE_ROOT / "document_bundle"
OUTLINE_PATH = FIXTURE_ROOT / "outline.json"
LAYOUT_MAP_PATH = PROJECT_ROOT / "templates/template_layout_map.json"
TEMPLATE_PATH = PROJECT_ROOT / "templates/financial_report_template_v1.pptx"


def _load(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def _write(path: Path, value: dict) -> None:
    path.write_text(
        json.dumps(value, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _profile(tmp_path: Path) -> Path:
    profile = build_template_profile(
        _load(LAYOUT_MAP_PATH),
        TEMPLATE_PATH,
        profile_id="week3-t3.7-test",
    )
    path = tmp_path / "template_profile.json"
    _write(path, profile)
    return path


def _run(*args: str) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        [sys.executable, str(PROJECT_ROOT / "main.py"), *args],
        cwd=PROJECT_ROOT,
        text=True,
        capture_output=True,
        check=False,
    )


def test_run_pipeline_publishes_complete_hashed_output(tmp_path):
    output = tmp_path / "pipeline-output"
    result = _run(
        "run-pipeline",
        str(BUNDLE_PATH),
        "--outline-input",
        str(OUTLINE_PATH),
        "--template-profile",
        str(_profile(tmp_path)),
        "--output-dir",
        str(output),
        "--candidate-mode",
        "active",
    )

    assert result.returncode == 0, result.stderr
    expected = {
        "document_bundle/document.json",
        "slide_outline.json",
        "numeric_fact_ledger.json",
        "metric_groups.json",
        "numeric_audit.json",
        "visualization_warnings.json",
        "template_profile.json",
        "visualizations/visualization_manifest.json",
        "compiled_layout_plan.json",
        "presentation.pptx",
        "run_manifest.json",
    }
    assert all((output / relative).is_file() for relative in expected)

    run_manifest = _load(output / "run_manifest.json")
    compiled_plan = _load(output / "compiled_layout_plan.json")
    audit = _load(output / "numeric_audit.json")
    metric_groups = _load(output / "metric_groups.json")
    assert run_manifest["status"] == "completed"
    assert audit["status"] == "passed"
    assert metric_groups["group_count"] == 2
    assert run_manifest["warnings"]["visualization_count"] == 0
    assert run_manifest["hashes"]["presentation_sha256"] == _sha256(
        output / "presentation.pptx"
    )
    assert compiled_plan["asset_root"] == str(
        (output / "document_bundle").resolve()
    )
    for artifact in run_manifest["artifacts"]:
        path = output / artifact["path"]
        assert path.is_file()
        assert artifact["sha256"] == _sha256(path)

    presentation = Presentation(output / "presentation.pptx")
    native_chart_count = sum(
        shape.has_chart
        for slide in presentation.slides
        for shape in slide.shapes
    )
    native_table_count = sum(
        shape.has_table
        for slide in presentation.slides
        for shape in slide.shapes
    )
    assert native_chart_count == 2
    assert native_table_count == 1


def test_missing_source_data_is_warning_but_explicit_verification_failure_blocks():
    outline = {
        "slides": [
            {
                "visual_candidates": [
                    {"candidate_id": "visual_explicit"}
                ]
            }
        ]
    }
    issues = [
        GenerationIssue(
            "slide_001",
            "cand_auto",
            "chart",
            "no_traceable_source_data",
        ),
        GenerationIssue(
            "slide_002",
            "visual_explicit",
            "table",
            "no_traceable_source_data",
        ),
        GenerationIssue(
            "slide_003",
            "visual_explicit",
            "chart",
            "verification_failed: source is outside the allowed evidence scope",
        ),
        GenerationIssue(
            "slide_004",
            "visual_explicit",
            "chart",
            "verification_failed: reject.mixed_metric: revenue and net_profit",
        ),
    ]

    blocking, warnings = _partition_generation_issues(outline, issues)

    assert [issue.visualization_id for issue in blocking] == [
        "visual_explicit"
    ]
    assert [issue.visualization_id for issue in warnings] == [
        "cand_auto",
        "visual_explicit",
        "visual_explicit",
    ]
    assert blocking[0].reason.startswith("verification_failed:")


def test_p0_failure_returns_nonzero_and_publishes_no_output(tmp_path):
    invalid_outline = deepcopy(_load(OUTLINE_PATH))
    invalid_outline["slides"][0]["evidence_refs"][0]["id"] = "missing-block"
    outline_path = tmp_path / "invalid-outline.json"
    _write(outline_path, invalid_outline)
    output = tmp_path / "failed-output"

    result = _run(
        "run-pipeline",
        str(BUNDLE_PATH),
        "--outline-input",
        str(outline_path),
        "--template-profile",
        str(_profile(tmp_path)),
        "--output-dir",
        str(output),
    )

    assert result.returncode != 0
    assert "pipeline stage=layout_preflight" in result.stderr
    assert not output.exists()
    assert not list(tmp_path.glob(".failed-output.staging-*"))
    failure = _load(tmp_path / "failed-output.failure" / "failure.json")
    assert failure["status"] == "failed"
    assert failure["failed_stage"] == "layout_preflight"
    assert failure["last_successful_stage"] == "candidate_locator"
    assert failure["published_pptx"] is False


def test_render_failure_does_not_publish_or_claim_a_presentation(tmp_path):
    wrong_template = tmp_path / "wrong-template.pptx"
    wrong_template.write_bytes(TEMPLATE_PATH.read_bytes() + b"tampered")
    output = tmp_path / "render-failed-output"

    result = _run(
        "run-pipeline",
        str(BUNDLE_PATH),
        "--outline-input",
        str(OUTLINE_PATH),
        "--template-profile",
        str(_profile(tmp_path)),
        "--template",
        str(wrong_template),
        "--output-dir",
        str(output),
    )

    assert result.returncode != 0
    assert "pipeline stage=render" in result.stderr
    assert "Pipeline completed:" not in result.stdout
    assert "Presentation:" not in result.stdout
    assert not output.exists()
    failure = _load(tmp_path / "render-failed-output.failure" / "failure.json")
    assert failure["failed_stage"] == "render"
    assert failure["published_pptx"] is False


def test_nonempty_output_directory_is_preserved(tmp_path):
    output = tmp_path / "existing-output"
    output.mkdir()
    sentinel = output / "keep.txt"
    sentinel.write_text("keep", encoding="utf-8")

    result = _run(
        "run-pipeline",
        str(BUNDLE_PATH),
        "--outline-input",
        str(OUTLINE_PATH),
        "--template-profile",
        str(_profile(tmp_path)),
        "--output-dir",
        str(output),
    )

    assert result.returncode != 0
    assert "pipeline stage=preflight" in result.stderr
    assert sentinel.read_text(encoding="utf-8") == "keep"
    assert not (output / "presentation.pptx").exists()

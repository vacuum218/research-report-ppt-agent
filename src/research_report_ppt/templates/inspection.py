#!/usr/bin/env python3
"""Inspect object names and structure in a PowerPoint template.

Example:
    python inspect_template.py financial_report_template_v1.pptx \
        --output template_objects.json

Requires:
    python-pptx>=0.6.23
"""

from __future__ import annotations

import re
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

try:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
except ImportError as exc:  # pragma: no cover - depends on local environment
    raise SystemExit(
        "缺少 python-pptx。请先执行：python -m pip install python-pptx"
    ) from exc


EMU_PER_INCH = 914400
GENERIC_NAME_PATTERN = re.compile(
    r"^(Chart|Picture|Image|TextBox|Text Box|Rectangle|Group|Table|Shape)(?:\s*\d+)?$",
    re.IGNORECASE,
)


def enum_name(value: Any) -> str:
    """Return a stable readable name for python-pptx enum values."""
    name = getattr(value, "name", None)
    if name:
        return str(name)
    text = str(value)
    return text.split(" ", 1)[0] if text else "UNKNOWN"


def inches(emu: int | None) -> float | None:
    if emu is None:
        return None
    return round(int(emu) / EMU_PER_INCH, 4)


def normalized(value: int | None, total: int) -> float | None:
    if value is None or not total:
        return None
    return round(int(value) / int(total), 6)


def clean_text(text: str, limit: int) -> str:
    compact = " ".join((text or "").split())
    if len(compact) <= limit:
        return compact
    return compact[: max(0, limit - 1)] + "…"


def safe_attr(obj: Any, name: str, default: Any = None) -> Any:
    try:
        return getattr(obj, name)
    except Exception:
        return default


def chart_info(shape: Any) -> dict[str, Any] | None:
    if not bool(safe_attr(shape, "has_chart", False)):
        return None

    chart = shape.chart
    result: dict[str, Any] = {
        "chart_type": enum_name(safe_attr(chart, "chart_type", "UNKNOWN")),
        "series_count": len(chart.series),
        "has_legend": bool(safe_attr(chart, "has_legend", False)),
        "has_title": bool(safe_attr(chart, "has_title", False)),
    }

    series = []
    for item in chart.series:
        series.append({"name": str(safe_attr(item, "name", ""))})
    result["series"] = series
    return result


def table_info(shape: Any) -> dict[str, Any] | None:
    if not bool(safe_attr(shape, "has_table", False)):
        return None
    table = shape.table
    return {
        "rows": len(table.rows),
        "columns": len(table.columns),
    }


def placeholder_info(shape: Any) -> dict[str, Any] | None:
    if not bool(safe_attr(shape, "is_placeholder", False)):
        return None
    fmt = shape.placeholder_format
    return {
        "idx": int(fmt.idx),
        "type": enum_name(fmt.type),
    }


def inspect_shape(
    shape: Any,
    slide_width: int,
    slide_height: int,
    text_preview_length: int,
    include_text: bool,
    path: str,
) -> dict[str, Any]:
    left = safe_attr(shape, "left")
    top = safe_attr(shape, "top")
    width = safe_attr(shape, "width")
    height = safe_attr(shape, "height")

    item: dict[str, Any] = {
        "path": path,
        "name": str(safe_attr(shape, "name", "")),
        "shape_id": int(safe_attr(shape, "shape_id", 0)),
        "shape_type": enum_name(safe_attr(shape, "shape_type", "UNKNOWN")),
        "position_emu": {
            "left": int(left) if left is not None else None,
            "top": int(top) if top is not None else None,
            "width": int(width) if width is not None else None,
            "height": int(height) if height is not None else None,
        },
        "position_in": {
            "left": inches(left),
            "top": inches(top),
            "width": inches(width),
            "height": inches(height),
        },
        "position_normalized": {
            "left": normalized(left, slide_width),
            "top": normalized(top, slide_height),
            "width": normalized(width, slide_width),
            "height": normalized(height, slide_height),
        },
        "rotation": float(safe_attr(shape, "rotation", 0.0) or 0.0),
        "has_text_frame": bool(safe_attr(shape, "has_text_frame", False)),
        "is_placeholder": bool(safe_attr(shape, "is_placeholder", False)),
        "has_chart": bool(safe_attr(shape, "has_chart", False)),
        "has_table": bool(safe_attr(shape, "has_table", False)),
    }

    if item["has_text_frame"]:
        text = safe_attr(shape, "text", "") or ""
        item["text_length"] = len(text)
        if include_text:
            item["text_preview"] = clean_text(text, text_preview_length)

    placeholder = placeholder_info(shape)
    if placeholder:
        item["placeholder"] = placeholder

    chart = chart_info(shape)
    if chart:
        item["chart"] = chart

    table_data = table_info(shape)
    if table_data:
        item["table"] = table_data

    if safe_attr(shape, "shape_type") == MSO_SHAPE_TYPE.GROUP:
        children = []
        for index, child in enumerate(shape.shapes, start=1):
            child_path = f"{path}.children[{index}]"
            children.append(
                inspect_shape(
                    child,
                    slide_width,
                    slide_height,
                    text_preview_length,
                    include_text,
                    child_path,
                )
            )
        item["children"] = children

    return item


def flatten_objects(objects: Iterable[dict[str, Any]]) -> Iterable[dict[str, Any]]:
    for item in objects:
        yield item
        yield from flatten_objects(item.get("children", []))


def inspect_presentation(
    pptx_path: Path,
    text_preview_length: int,
    include_text: bool,
) -> dict[str, Any]:
    prs = Presentation(str(pptx_path))
    slide_width = int(prs.slide_width)
    slide_height = int(prs.slide_height)

    slides: list[dict[str, Any]] = []
    total_types: Counter[str] = Counter()

    for slide_number, slide in enumerate(prs.slides, start=1):
        objects = []
        for object_number, shape in enumerate(slide.shapes, start=1):
            objects.append(
                inspect_shape(
                    shape,
                    slide_width,
                    slide_height,
                    text_preview_length,
                    include_text,
                    f"slides[{slide_number}].objects[{object_number}]",
                )
            )

        flattened = list(flatten_objects(objects))
        type_counts = Counter(item["shape_type"] for item in flattened)
        name_counts = Counter(item["name"] for item in flattened if item["name"])
        duplicate_names = sorted(
            name for name, count in name_counts.items() if count > 1
        )
        generic_names = sorted(
            {item["name"] for item in flattened if GENERIC_NAME_PATTERN.match(item["name"])}
        )
        total_types.update(type_counts)

        slides.append(
            {
                "slide_number": slide_number,
                "object_count": len(flattened),
                "top_level_object_count": len(objects),
                "object_type_counts": dict(sorted(type_counts.items())),
                "object_names": [item["name"] for item in flattened],
                "duplicate_object_names": duplicate_names,
                "generic_object_names": generic_names,
                "objects": objects,
            }
        )

    return {
        "schema_version": "1.0",
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "source_file": pptx_path.name,
        "source_path": str(pptx_path.resolve()),
        "presentation": {
            "slide_count": len(prs.slides),
            "slide_width_emu": slide_width,
            "slide_height_emu": slide_height,
            "slide_width_in": inches(slide_width),
            "slide_height_in": inches(slide_height),
            "object_type_counts": dict(sorted(total_types.items())),
        },
        "slides": slides,
    }


def print_summary(data: dict[str, Any]) -> None:
    presentation = data["presentation"]
    print(
        f"文件：{data['source_file']} | "
        f"页数：{presentation['slide_count']} | "
        f"尺寸：{presentation['slide_width_in']} × "
        f"{presentation['slide_height_in']} 英寸"
    )

    for slide in data["slides"]:
        print(f"\nSlide {slide['slide_number']:02d} ({slide['object_count']} objects)")
        for name in slide["object_names"]:
            print(f"  - {name}")
        if slide["duplicate_object_names"]:
            names = ", ".join(slide["duplicate_object_names"])
            print(f"  [警告] 本页存在重复对象名：{names}")
        if slide["generic_object_names"]:
            names = ", ".join(slide["generic_object_names"])
            print(f"  [提示] 建议重命名通用对象名：{names}")
